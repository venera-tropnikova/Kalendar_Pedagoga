"""Real-document gate. Filename knowledge is confined to the corpus manifest."""
from __future__ import annotations
from collections import Counter
from dataclasses import asdict
import difflib
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import subprocess
import sys
from tempfile import TemporaryDirectory
import xml.etree.ElementTree as ET

from docx import Document
from docx.table import Table as DocxTable

from calendar_pedagoga.lossless_document import extract_document, extract_bytes, LibreOfficeConverter
from calendar_pedagoga.lossless_document.formats import _PROFILE, office_environment
from .audit import audit, inventory, structure_digest, W

ODT_TEXT = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
ODT_TABLE = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
ODT_OFFICE = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'


def corpus_paths(manifest):
    repo = Path(__file__).resolve().parents[2]
    downloads = Path(os.environ.get('KP_CORPUS_DOWNLOADS', Path.home() / 'Downloads'))
    records = json.loads(Path(manifest).read_text(encoding='utf-8'))
    for record in records:
        record['path'] = (repo / 'references' if record['location'] == 'references' else downloads) / record['filename']
        assert record['path'].is_file(), f"Required real corpus document missing: {record['path']}"
        assert sha256(record['path'].read_bytes()).hexdigest() == record['sha256'], f"Corpus bytes changed: {record['path']}"
    return records


def verify_docx_merge_view(document):
    """Second table implementation: python-docx visual grid, main story only."""
    doc = Document(BytesIO(document.package_bytes))
    raw_tables = list(doc.element.iter(W + 'tbl'))
    extracted_tables = [t for t in document.tables if t.source.part_uri == '/word/document.xml']
    assert len(raw_tables) == len(extracted_tables)
    for raw, extracted in zip(raw_tables, extracted_tables):
        visual = DocxTable(raw, doc._body)
        groups = {}
        for row_number, row in enumerate(visual.rows):
            offset = row.grid_cols_before
            for column, cell in enumerate(row.cells, offset):
                groups.setdefault(cell._tc, []).append((row_number, column))
        oracle = sorted(tuple(value) for value in groups.values())
        assert oracle == sorted(c.grid_slots for c in extracted.logical_cells), 'Merged grid differs from python-docx'
    return {'tables_checked':len(raw_tables),'status':'PASS'}


def fodt_witness(source_bytes, converter, output):
    with TemporaryDirectory(prefix='kp_a_fodt_') as temp:
        root=Path(temp); profile=root/'profile'; (profile/'user').mkdir(parents=True)
        (profile/'user/registrymodifications.xcu').write_text(_PROFILE,encoding='utf-8')
        source=root/'input.doc'; source.write_bytes(source_bytes)
        result=subprocess.run([str(converter.executable),f'-env:UserInstallation={profile.as_uri()}',
                               '--headless','--nologo','--nodefault','--norestore','--convert-to','fodt',
                               '--outdir',str(root),str(source)],capture_output=True,timeout=90,check=True,cwd=root,env=office_environment())
        data=(root/'input.fodt').read_bytes()
        output.write_bytes(data)
        return {'sha256':sha256(data).hexdigest(),'stderr':result.stderr.decode('utf-8','replace')}


def compare_fodt(document, path):
    root=ET.fromstring(path.read_bytes()); body=root.find(ODT_OFFICE+'body')
    def normalized(text): return re.sub(r'\s+',' ',text).strip()
    original=[''.join(e.itertext()) for e in body.iter() if e.tag in (ODT_TEXT+'p',ODT_TEXT+'h')]
    _, xml_nodes, _=inventory(document.package_bytes)
    paragraphs=[(a,e) for a,e in xml_nodes if a[0]=='/word/document.xml' and e.tag==W+'p']
    converted=[''.join(t.text or '' for t in p.iter(W+'t')) for _,p in paragraphs]
    a=list(map(normalized,original)); b=list(map(normalized,converted))
    assert [s for s in a if s]==[s for s in b if s], 'DOC conversion changed nonempty paragraph order/text'
    parent_map={c:p for _,p in xml_nodes for c in p}
    explained=[]
    for op,i,j,k,l in difflib.SequenceMatcher(a=a,b=b,autojunk=False).get_opcodes():
        if op=='equal':continue
        assert op=='insert' and all(not text for text in b[k:l]), (op,i,j,k,l)
        for index in range(k,l):
            address, p=paragraphs[index]; parent=parent_map[p]
            merge=parent.find(W+'tcPr/'+W+'vMerge')
            if parent.tag==W+'tc' and merge is not None and merge.get(W+'val','continue')=='continue':
                reason='empty paragraph required in vertical merge continuation cell'
            elif p.find(W+'pPr/'+W+'sectPr') is not None:
                reason='empty paragraph carrying a section break'
            else:
                raise AssertionError(f'Unexplained added empty paragraph at {address}')
            explained.append({'part':address[0],'element_path':address[1],'paragraph_index':index,'reason':reason})
    cells=list(body.iter(ODT_TABLE+'table-cell'))
    table_count=len(list(body.iter(ODT_TABLE+'table')))
    merged=sum(int(c.get(ODT_TABLE+'number-columns-spanned','1'))>1 or int(c.get(ODT_TABLE+'number-rows-spanned','1'))>1 for c in cells)
    assert table_count==len(document.tables)
    assert len(cells)==sum(len(t.logical_cells) for t in document.tables)
    assert merged==document.counts()['merges']
    return {'nonempty_paragraphs':len([s for s in a if s]),'ordered_nonempty_text_equal':True,
            'fodt_paragraphs':len(a),'docx_main_paragraphs':len(b),'explained_empty_paragraphs':explained,
            'logical_cells':len(cells),'merges':merged,'tables':table_count}


def word_blocks(text):
    # Word's paragraph, cell-end and section-break delimiters are distinct source controls.
    return [re.sub(r'\s+',' ',p.replace('\x07','').replace('\x01','')).strip()
            for p in re.split('[\r\x0c]',text) if p.replace('\x07','').replace('\x01','').strip()]


def verify_corpus(manifest, artifact_dir):
    artifact_dir=Path(artifact_dir); artifact_dir.mkdir(parents=True,exist_ok=True)
    records=corpus_paths(manifest)
    converter=LibreOfficeConverter.discover()
    results=[]; word_inputs=[]
    for record in records:
        key=record['id']; path=record['path']
        print('EXTRACT',key,flush=True)
        first=extract_document(path,converter=converter)
        checks=audit(first)
        checks['merged_grid_oracle']=verify_docx_merge_view(first)
        repeated=extract_document(path,converter=converter)
        assert first.nodes==repeated.nodes and first.paragraphs==repeated.paragraphs and first.tables==repeated.tables
        assert first.year_candidates==repeated.year_candidates
        with TemporaryDirectory(prefix='kp_a_rename_') as temp:
            renamed=Path(temp)/'arbitrary-name.wrong-extension'
            renamed.write_bytes(path.read_bytes())
            third=extract_document(renamed,converter=converter)
        assert first.nodes==third.nodes and first.paragraphs==third.paragraphs and first.tables==third.tables
        assert first.year_candidates==third.year_candidates
        checks.update({'repeat_ids_and_structure':True,'renamed_ids_and_structure':True,'no_year_assignment':True})
        converted=extract_bytes(first.package_bytes)
        audit(converted)
        assert structure_digest(first)==structure_digest(converted)
        assert first.counts()==converted.counts()
        converted_path=artifact_dir/(key+'.docx');converted_path.write_bytes(first.package_bytes)
        result={'id':key,'label':record['label'],'filename':record['filename'],'format':first.source_format,
                'sha256':first.source_sha256,'package_sha256':first.package_sha256,'counts':first.counts(),
                'checks':checks,'diagnostics':[asdict(d) for d in first.diagnostics],
                'conversion_log':[asdict(e) for e in first.conversion_log],
                'losses':[],'unexplained_duplicates':[],'ambiguities':[]}
        if first.alternative_branches:
            result['ambiguities'].append({'kind':'alternative_XML_representations','branches':len(first.alternative_branches),
                                          'treatment':'Choice and Fallback retained as distinct branches; neither selected nor flattened'})
        if first.source_format=='DOC':
            copied=artifact_dir/(key+'.doc');copied.write_bytes(path.read_bytes())
            fodt=artifact_dir/(key+'.fodt')
            result['fodt_witness']=fodt_witness(path.read_bytes(),converter,fodt)
            result['doc_conversion_comparison']=compare_fodt(first,fodt)
            result['ambiguities'].append({'kind':'conversion_provenance','treatment':'SourceSpan addresses converted OOXML; original DOC bytes and both hashes retained'})
            word_inputs.extend([{'key':key,'format':'DOC','path':str(copied)},
                                {'key':key,'format':'DOCX','path':str(converted_path)}])
        results.append(result)
    inputs=artifact_dir/'word_inputs.json';inputs.write_text(json.dumps(word_inputs,ensure_ascii=False),encoding='utf-8')
    witness=artifact_dir/'word_witness.json'
    subprocess.run([sys.executable,str(Path(__file__).with_name('word_witness.py')),str(witness),str(inputs)],check=True,timeout=180)
    snapshots=json.loads(witness.read_text(encoding='utf-8'))
    for original, converted in zip(snapshots[::2],snapshots[1::2]):
        a=word_blocks(next(s['text'] for s in original['stories'] if s['type']==1))
        b=word_blocks(next(s['text'] for s in converted['stories'] if s['type']==1))
        assert a==b, f"Independent Word importer disagrees: {original['key']}"
        result=next(r for r in results if r['id']==original['key'])
        assert len(a)==result['doc_conversion_comparison']['nonempty_paragraphs']
        assert original['main_table_count']==converted['main_table_count']==result['counts']['tables']
        result['word_witness']={'version':original['word_version'],'ordered_nonempty_blocks':len(a),
                                'doc_and_docx_equal':True,'tables':original['main_table_count'],
                                'scope':'main story; auxiliary stories and compatibility branches inventoried separately in OOXML',
                                'source_snapshot_sha256':sha256(json.dumps(original,ensure_ascii=False,sort_keys=True).encode()).hexdigest()}
        # Non-content separator stories can be synthesized/omitted by Word/LibreOffice.
        auxiliary=[]
        for snapshot in (original,converted):
            auxiliary.append([{'type':s['type'],'paragraph_count':s['paragraph_count'],
                               'text_sha256':sha256(s['text'].encode()).hexdigest()}
                              for s in snapshot['stories'] if s['type']!=1])
        result['word_auxiliary_stories']={'DOC':auxiliary[0],'DOCX':auxiliary[1]}
        package=(artifact_dir/(original['key']+'.docx')).read_bytes()
        extracted=extract_bytes(package)
        aux_paragraphs=[p for p in extracted.paragraphs if p.source.part_uri!='/word/document.xml']
        aux_codes={re.sub(r'\s+',' ',f.raw_text).strip() for f in extracted.fragments
                   if f.role=='instrText' and f.source.part_uri!='/word/document.xml'}
        explain_aux=[]
        for story in original['stories']:
            if story['type']==1:continue
            text=re.sub(r'\s+',' ',story['text']).strip()
            if not text:continue
            if story['type'] in (12,13,15,16) and not text.strip('\x03\x04'):
                explain_aux.append({'story_type':story['type'],'reason':'automatic footnote/endnote separator control; no content text'})
                continue
            fields=story.get('fields',[])
            field_result=' '.join(f['result'].strip() for f in fields).strip()
            if fields and text==field_result:
                codes=[re.sub(r'\s+',' ',f['code']).strip() for f in fields]
                assert all(code in aux_codes for code in codes), 'Lost auxiliary field instruction'
                explain_aux.append({'story_type':story['type'],'reason':'field instruction preserved; cached display may change and compatibility branches remain separate',
                                    'field_codes':codes,'original_cached_text':text})
            else:
                assert text in [re.sub(r'\s+',' ',p.raw_text).strip() for p in aux_paragraphs], 'Unexplained auxiliary text loss'
                explain_aux.append({'story_type':story['type'],'reason':'literal auxiliary text present in converted XML'})
        result['doc_auxiliary_explanations']=explain_aux
    assert len(results)==5 and len({r['sha256'] for r in results})==5
    return {'schema':'lossless-extraction-corpus-report/1','base_commit':'0f49496dc06e361b068d5ebc048f2028da506bde',
            'status':'PASS','documents':results,
            'limitations':['No visual/layout equivalence assertion; Stage A preserves source evidence.',
                           'Native Word verifies all nonempty main-story blocks; FODT verifies text order and merged-cell counts.',
                           'DOC conversion may expand linked footers into multiple parts and change PAGE cached values; original bytes remain available.',
                           'No year assignment, plan interpretation, matching, scheduling or calendar export performed.']}
