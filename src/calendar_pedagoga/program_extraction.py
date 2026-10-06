"""Ordered source blocks and explicit content-boundary candidates for program import."""
from dataclasses import dataclass
from io import BytesIO
import re
import os
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


class ProgramReadError(ValueError):
    """A source location cannot be read safely."""


@dataclass(frozen=True)
class TextBlock:
    text: str
    location: str
    page: int | None = None
    bold: bool = False
    alignment: int | None = None
    table: str | None = None
    heading_hint: bool = False

    @property
    def raw_text(self):
        return self.text

    @property
    def normalized_text(self):
        return re.sub(r"\s+", " ", self.text).strip()

    @property
    def runs(self):
        # Existing metadata/heading readers need only text and bold.
        return (self,)


@dataclass(frozen=True)
class ContentRegion:
    id: str
    start: int
    end: int
    year: int | None
    label: str
    evidence: str
    needs_confirmation: bool = False
    heading_indices: tuple[int, ...] = ()


@dataclass(frozen=True)
class ProgramDocument:
    paragraphs: tuple[TextBlock, ...]

    @property
    def source_labels(self):
        return frozenset(label_key(b.text) for b in self.paragraphs
                         if b.heading_hint and b.text.strip() and not b.text.strip().isdigit())


def label_key(text):
    return re.sub(r"\s+", " ", re.sub(r"[.\s]*\(\s*\d+(?:[.,]\d+)?\s*\)\s*$", "", text)).strip(" .:").casefold()


def read_docx(data):
    try:
        doc = Document(BytesIO(data))
    except Exception as error:
        raise ProgramReadError("DOCX повреждён: не удалось прочитать структуру документа.") from error
    blocks = []
    title_paragraphs = set()
    for table in doc.element.body.iter(qn("w:tbl")):
        rows = table.findall(qn("w:tr"))
        columns = set()
        header_rows = set()
        for row in rows[:3]:
            for index, cell in enumerate(row.findall(qn("w:tc"))):
                text = " ".join(t.text or "" for t in cell.iter(qn("w:t")))
                if re.search(r"^(?:наименование|название)\s+(?:раздел|тем)|^(?:тема|раздел)(?:\s+и\s+тема)?$", text.strip(), re.I):
                    columns.add(index)
                    header_rows.add(row)
        for row in rows:
            if row in header_rows:
                continue
            for index, cell in enumerate(row.findall(qn("w:tc"))):
                if index in columns:
                    title_paragraphs.update(cell.iter(qn("w:p")))

    def walk(parent, path, table=None):
        for index, child in enumerate(parent):
            here = f"{path}/{child.tag.rsplit('}', 1)[-1]}[{index + 1}]"
            if child.tag == qn("w:p"):
                p = Paragraph(child, doc)
                runs = [r for r in p.runs if r.text.strip()]
                blocks.append(TextBlock(p.text, here, bold=bool(runs) and all(r.bold is True for r in runs),
                                        alignment=1 if child.xpath("./w:pPr/w:jc[@w:val='center']") else None, table=table, heading_hint=child in title_paragraphs))
            elif child.tag in (qn("w:tbl"), qn("w:tr"), qn("w:tc"), qn("w:sdt"), qn("w:sdtContent")):
                # Physical XML cells, not row.cells: merged cells are visited once.
                walk(child, here, here if child.tag == qn("w:tbl") else table)
    walk(doc.element.body, "document/body")
    return ProgramDocument(tuple(blocks))


def ocr_tessdata():
    """Resolve language data at runtime, including ordinary app launches."""
    import pymupdf
    candidates = []
    if os.environ.get("TESSDATA_PREFIX"):
        candidates.append(Path(os.environ["TESSDATA_PREFIX"]))
    if os.environ.get("LOCALAPPDATA"):
        candidates.append(Path(os.environ["LOCALAPPDATA"]) / "CalendarPedagoga" / "tessdata")
    candidates.extend([Path("/usr/share/tesseract-ocr/5/tessdata"),
                       Path("C:/Program Files/Tesseract-OCR/tessdata")])
    for candidate in candidates:
        if (candidate / "rus.traineddata").is_file():
            return str(candidate)
    try:
        candidate = Path(pymupdf.get_tessdata())
        if (candidate / "rus.traineddata").is_file():
            return str(candidate)
    except (RuntimeError, OSError):
        pass
    raise ProgramReadError("OCR недоступен: не установлена русская модель rus.traineddata. "
                           "Установите tesseract-ocr-rus или задайте TESSDATA_PREFIX.")


def _ocr_page(page):
    tessdata = ocr_tessdata()
    try:
        return page.get_textpage_ocr(language="rus", dpi=300, full=True, tessdata=tessdata)
    except Exception as error:
        raise ProgramReadError("OCR не выполнен: движок или русская модель недоступны "
                               f"({type(error).__name__}).") from error


def read_pdf(data):
    import pymupdf
    blocks, unreadable = [], []
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ProgramReadError("PDF защищён паролем. Загрузите доступную для чтения копию.")
            for number, page in enumerate(doc, 1):
                page_blocks = []
                extracted = page.get_text("dict", sort=True)
                native_letters = sum(c.isalpha() for block in extracted["blocks"] if block["type"] == 0
                                     for line in block["lines"] for span in line["spans"] for c in span["text"])
                used_ocr = native_letters < 3 or (page.get_images() and native_letters < 80)
                if used_ocr:
                    try:
                        textpage = _ocr_page(page)
                        extracted = page.get_text("dict", textpage=textpage, sort=True)
                    except ProgramReadError as error:
                        unreadable.append((number, str(error)))
                        continue
                for bi, block in enumerate(extracted["blocks"]):
                    if block["type"] != 0:
                        continue
                    for li, line in enumerate(block["lines"]):
                        text = "".join(span["text"] for span in line["spans"])
                        if text.strip():
                            page_blocks.append(TextBlock(text, f"страница {number}, {'OCR, ' if used_ocr else ''}блок {bi + 1}, строка {li + 1}",
                                                         page=number))
                letters = sum(c.isalpha() for b in page_blocks for c in b.text)
                if letters < 3:
                    unreadable.append((number, "OCR завершён, но читаемый текст не найден."))
                blocks.extend(page_blocks)
    except ProgramReadError:
        raise
    except Exception as error:
        raise ProgramReadError("PDF повреждён или не может быть прочитан.") from error
    if unreadable:
        raise ProgramReadError("Не удалось распознать текст PDF: страницы "
                               + ", ".join(str(n) for n, _ in unreadable) + ". "
                               + "; ".join(f"Страница {n}: {reason}" for n, reason in unreadable)
                               + " Проверьте качество скана и OCR на русском; загрузите читаемую копию.")
    return ProgramDocument(tuple(blocks))


CONTENT = re.compile(r"^(?:\d+(?:\.\d+)*[.)]?\s*)?(?:содержание(?:\s+(?:образовательной|учебной|рабочей))?\s+"
                     r"(?:программы|обучения|курса|учебного\s+(?:плана|материала)|учебно[-–]тематического\s+плана)|содержательная\s+часть)\b", re.I)
STOP = re.compile(r"^(?:(?:ожидаемые|планируемые)\s+результаты|по окончани[юя]|"
                  r"должны\s+(?:знать|уметь)|методическое\s+обеспечение|"
                  r"список\s+литературы|условия\s+реализации)", re.I)
YEAR = re.compile(r"\b(\d+|первый|первого|второй|второго|третий|третьего|четвертый|четвёртый|четвертого|четвёртого|I{1,3}|IV)"
                  r"\s*(?:[-–]?(?:й|ый|ого|го))?\s+год(?:а)?\s+обучения\b", re.I)


def year_of(text):
    match = YEAR.search(text)
    if not match:
        return None
    token = match[1].casefold()
    if token.isdigit():
        return int(token)
    for prefix, value in (("перв", 1), ("втор", 2), ("трет", 3), ("треть", 3), ("четв", 4)):
        if token.startswith(prefix):
            return value
    return {"i": 1, "ii": 2, "iii": 3, "iv": 4}.get(token)


def standalone_year(text):
    return bool(YEAR.fullmatch(text.strip(" ():.")))


def is_heading(block, labels):
    text = block.text.strip()
    if not text or re.match(r"^(?:теория|практика|практическ\w+\s+работ\w*)[.:]?$", text, re.I):
        return False
    return bool(re.match(r"^(?:раздел|тема)\s+\d+[.:\s]", text, re.I)
                or re.match(r"^\d+(?:\.\d+)*[.)]?\s+\S", text)
                or (block.bold and len(text) < 180)
                or label_key(text) in labels)


def content_heading(text):
    text = re.sub(r"\s+", " ", text).strip()
    match = CONTENT.match(text.strip())
    if not match:
        return False
    tail = text.strip()[match.end():].strip(" .:()")
    tail = re.sub(r'^[«"][^»"]+[»"]', "", tail).strip(" .:()")
    return not tail or standalone_year(tail)


def content_regions(document):
    blocks = document.paragraphs
    starts = [(i + 1, year_of(b.text), False) for i, b in enumerate(blocks) if content_heading(b.text)]
    if not starts:
        # Source-derived repeated labels locate narrative following tables; no topic dictionary.
        labels = document.source_labels
        for i, b in enumerate(blocks):
            if not b.table and label_key(b.text) in labels and b.text.strip():
                if any(x.table for x in blocks[:i]):
                    starts.append((i, year_of(b.text), True))
                    break
        if not starts:
            numbered = [i for i, b in enumerate(blocks) if not b.table
                        and re.match(r"^(?:раздел|тема)\s+\d+[.:\s]", b.text.strip(), re.I)]
            if numbered:
                starts = [(numbered[0], None, True)]
    regions = []
    for pos, (start, year, inferred) in enumerate(starts):
        end = starts[pos + 1][0] - 1 if pos + 1 < len(starts) else len(blocks)
        markers = [(i, year_of(blocks[i].text)) for i in range(start, end)
                   if standalone_year(blocks[i].text)]
        scope_markers = []
        for index, marker_year in markers:
            previous_index, previous_year = scope_markers[-1] if scope_markers else (start, year)
            # Repeating the same year after outcomes does not reopen teaching content.
            if marker_year == previous_year and any(
                STOP.match(blocks[j].text.strip()) for j in range(previous_index, index)
            ):
                continue
            scope_markers.append((index, marker_year))
        markers = scope_markers
        intervals = ([(i + 1, markers[n + 1][0] if n + 1 < len(markers) else end, y)
                      for n, (i, y) in enumerate(markers)] if markers else [(start, end, year)])
        for lo, hi, scope in intervals:
            hi = next((i for i in range(lo, hi) if STOP.match(blocks[i].text.strip())), hi)
            if not any(b.text.strip() for b in blocks[lo:hi]):
                continue
            heading_indices = ()
            if not any(is_heading(b, document.source_labels) for b in blocks[lo:hi]):
                # Short isolated lines are hypotheses only; explicit teacher confirmation is required.
                heading_indices = tuple(i for i in range(lo, hi - 1)
                                        if blocks[i].text.strip() and len(blocks[i].text.strip()) <= 100
                                        and not re.search(r"[.!?;:]", blocks[i].text)
                                        and blocks[i + 1].text.strip()
                                        and (len(blocks[i + 1].text) > len(blocks[i].text)
                                             or blocks[i + 1].text.rstrip().endswith(".")))
            location = next(b.text.strip() for b in blocks[lo:hi] if b.text.strip())[:100]
            regions.append(ContentRegion(f"{lo}:{hi}", lo, hi, scope,
                                         (f"{scope}-й год; " if scope else "") + location,
                                         "повтор названий из таблицы / последовательность разделов" if inferred
                                         else "заголовок содержания и порядок исходных блоков",
                                         inferred or bool(heading_indices), heading_indices))
    return tuple(regions)
