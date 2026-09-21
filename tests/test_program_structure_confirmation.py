from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

import pytest
from docx import Document

from calendar_pedagoga.confirmed_study_plan import (
    ConfirmedStudyPlan,
    confirmed_plan_from_external_utp,
    confirmed_plan_from_manual_rows,
)
from calendar_pedagoga.content_engine_v2 import LessonContentV2Row, build_lesson_content_v2
from calendar_pedagoga.content_generation import CalendarContentRow, build_content_model
from calendar_pedagoga.match_review import apply_match_reviews
from calendar_pedagoga.matching import ContentMatch, MatchStatus, match_utp_to_program
from calendar_pedagoga.parsing import Hours, Section, Topic, UtpMetadata, UtpParseResult, parse_utp
from calendar_pedagoga.production_readiness import (
    EMPTY_RESULT,
    SOURCE_NOT_MATCHED,
    production_readiness_codes,
)
from calendar_pedagoga.program_parsing import ProgramContentItem, ProgramData, parse_program
from calendar_pedagoga.program_structure_confirmation import (
    CONTENT_ORIGIN_EXCERPT,
    CONTENT_ORIGIN_MANUAL,
    DISPOSITION_EXCLUDED,
    DISPOSITION_MAPPED,
    DISPOSITION_UNRESOLVED,
    EXCLUSION_HEADING,
    EXCLUSION_OTHER_YEAR,
    EXCLUSION_SERVICE,
    PROVEN_MATCH_STATUSES,
    TOPIC_STATUS_DRAFT_READY,
    TOPIC_STATUS_UNRESOLVED,
    TOPIC_STATUS_USER_CONFIRMED,
    ProgramStructureConfirmationError,
    catalog_source_items,
    confirmation_is_current,
    confirm_program_structure,
    draft_source_ledger,
    draft_structure_rows,
    embedded_utp_candidates,
    file_digest,
    merge_source_ledger,
    needs_structure_confirmation,
    overlay_confirmed_program,
    overlay_unresolved_topic_edits,
    schedule_topics_from_candidate,
    select_embedded_utp,
    source_disposition_counts,
    structure_confirmation_scope,
    unmatched_source_topics,
    unresolved_schedule_rows,
    STRUCTURE_CONFIRM_BUTTON,
    STRUCTURE_EXCERPT_LABEL,
    STRUCTURE_ORIGIN_EXCERPT_LABEL,
    STRUCTURE_PRACTICE_CONTENT_LABEL,
    STRUCTURE_SOURCE_ITEM_LABEL,
    STRUCTURE_THEORY_CONTENT_LABEL,
)
from calendar_pedagoga.scheduling import build_schedule
from calendar_pedagoga.resolve_utp import resolve_utp
from calendar_pedagoga.upload_validation import UploadPurpose, ValidatedUpload, validate_upload


REFERENCES = Path(__file__).resolve().parents[1] / "references"
EXTRA_REFERENCES = Path(r"D:\Kalendar_Pedagoga\references")


def _corpus(*needles: str) -> Path | None:
    lowered = tuple(needle.casefold() for needle in needles)
    for root in (REFERENCES, EXTRA_REFERENCES):
        if not root.exists():
            continue
        matches = [
            path
            for path in root.iterdir()
            if all(needle in path.name.casefold() for needle in lowered)
        ]
        if matches:
            return matches[0]
    return None


def _program_data(*items: ProgramContentItem) -> ProgramData:
    return ProgramData(
        title="Неизвестная программа",
        duration=None,
        student_age=None,
        goal="Цель",
        tasks=("Задача",),
        lesson_forms=(),
        teaching_methods=(),
        expected_results=(),
        knowledge_outcomes=(),
        skill_outcomes=(),
        content_items=items,
    )


def _item(title: str, content: str = "Лепка из глины.") -> ProgramContentItem:
    return ProgramContentItem("1", title, content, title, 1)


def _manual_row(topic: str = "Лепка", theory: str = "1", practice: str = "1") -> dict[str, str]:
    return {
        "topic": topic,
        "theory_hours": theory,
        "practice_hours": practice,
        "theory_content": "Знакомство с материалом.",
        "practice_content": "Лепка из глины.",
        "source": "program",
        "match_status": MatchStatus.NOT_MATCHED.value,
    }


def _ledger_row(
    record,
    *,
    disposition: str = DISPOSITION_UNRESOLVED,
    exclusion_reason: str = "",
    mapped_topic: str = "",
) -> dict[str, str]:
    return {
        **record.as_dict(),
        "disposition": disposition,
        "exclusion_reason": exclusion_reason,
        "mapped_topic": mapped_topic,
    }


def _holdout_exclusion_reason(title: str, section: str) -> str | None:
    blob = f"{title}\n{section}"
    year1 = title in {"I год обучения", "Особенности работы 1 года обучения"}
    if year1:
        return None
    if any(
        token in blob
        for token in (
            "II год",
            "III год",
            "II года",
            "III года",
            "2 года",
            "3 года",
        )
    ):
        return EXCLUSION_OTHER_YEAR
    if title.startswith("Учебно-тематический план"):
        return EXCLUSION_HEADING
    return EXCLUSION_SERVICE


def _year_source_blocks(text: str) -> tuple[str, ...]:
    return tuple(
        block.strip()
        for block in re.split(r"(?=Тема №\s*\d+)", text or "")
        if block.strip().startswith("Тема №")
    )


def _holdout_excerpt_for_topic(source_text: str, topic_title: str) -> str:
    markers = (
        ("Вводное занятие", "Тема № 1 "),
        ("Аппликация", "Тема № 2 "),
        ("Скульптурная композиция", "Тема № 3 "),
        ("Декоративные украшения", "Тема № 4 "),
        ("Плетение", "Тема № 5 "),
        ("Солёное тесто", "Тема № 6 "),
    )
    for needle, prefix in markers:
        if needle.casefold() not in topic_title.casefold():
            continue
        for block in _year_source_blocks(source_text):
            if block.startswith(prefix) or needle.casefold() in block[:80].casefold():
                return block
    return ""


def _holdout_topic_rows(draft, source_text: str, source_item_id: str):
    rows = []
    for row in draft:
        excerpt = _holdout_excerpt_for_topic(source_text, row["topic"])
        if not excerpt:
            rows.append(
                {
                    **row,
                    "source_item_id": "",
                    "excerpt": "",
                    "content_origin": "",
                    "theory_content": "",
                    "practice_content": "",
                }
            )
            continue
        rows.append(
            {
                **row,
                "source_item_id": source_item_id,
                "excerpt": excerpt,
                "content_origin": CONTENT_ORIGIN_EXCERPT,
                "theory_content": "",
                "practice_content": "",
            }
        )
    return tuple(rows)


def _resolve_holdout_ledger(program, schedule_titles: tuple[str, ...]):
    del schedule_titles
    catalog = catalog_source_items(program.content_items)
    ledger = []
    mapped_rows = []
    for record in catalog:
        reason = _holdout_exclusion_reason(record.title, record.section or "")
        if record.title == "I год обучения":
            mapped_rows.append(record)
            ledger.append(
                _ledger_row(
                    record,
                    disposition=DISPOSITION_UNRESOLVED,
                )
            )
            continue
        if reason is None:
            ledger.append(
                _ledger_row(
                    record,
                    disposition=DISPOSITION_EXCLUDED,
                    exclusion_reason=EXCLUSION_HEADING,
                )
            )
            continue
        ledger.append(
            _ledger_row(
                record,
                disposition=DISPOSITION_EXCLUDED,
                exclusion_reason=reason,
            )
        )
    return catalog, tuple(ledger), tuple(mapped_rows)


def _program_docx_with_utp_years(*years: int) -> bytes:
    document = Document()
    document.add_paragraph("Дополнительная общеобразовательная программа «ТЕСТ»")
    document.add_paragraph("Цель: научить основам.")
    document.add_paragraph("Задачи: освоить приёмы.")
    document.add_paragraph("Содержание программы 1-го года обучения")
    document.add_paragraph("1. Введение")
    document.add_paragraph("Знакомство с программой.")
    roman = {1: "I", 2: "II", 3: "III"}
    for year in years:
        document.add_paragraph(
            f"Учебно-тематический план {roman.get(year, year)} год обучения"
        )
        table = document.add_table(rows=6, cols=5)
        for cell, value in zip(
            table.rows[0].cells,
            ("№", "Тема", "всего", "теория", "практика"),
            strict=True,
        ):
            cell.text = value
        rows = (
            ("1", f"Раздел {year}", "36", "12", "24"),
            ("1.1", f"Тема {year}.1", "12", "4", "8"),
            ("1.2", f"Тема {year}.2", "12", "4", "8"),
            ("1.3", f"Тема {year}.3", "12", "4", "8"),
            ("Итого", "", "36", "12", "24"),
        )
        for index, values in enumerate(rows, start=1):
            for cell, value in zip(table.rows[index].cells, values, strict=True):
                cell.text = value
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _card_program_docx() -> bytes:
    document = Document()
    document.add_paragraph("Дополнительная общеобразовательная программа «Карточки»")
    document.add_paragraph("Цель: научить основам.")
    document.add_paragraph("Задачи: освоить приёмы.")
    document.add_paragraph("Содержание программы 1-го года обучения")
    heading = document.add_paragraph()
    run = heading.add_run("I год обучения")
    run.bold = True
    document.add_paragraph(
        "Тема № 1 Введение в материал и правила безопасности. "
        "Знакомство с инструментами."
    )
    document.add_paragraph("Учебно-тематический план I год обучения")
    table = document.add_table(rows=5, cols=5)
    for cell, value in zip(
        table.rows[0].cells,
        ("№", "Тема", "всего", "теория", "практика"),
        strict=True,
    ):
        cell.text = value
    rows = (
        ("1", "Введение", "1", "1", "0"),
        ("2", "Итоговое занятие", "2", "2", "0"),
        ("3", "Конкурсы и выставки", "2", "0", "2"),
        ("Итого", "", "5", "3", "2"),
    )
    for index, values in enumerate(rows, start=1):
        for cell, value in zip(table.rows[index].cells, values, strict=True):
            cell.text = value
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _upload_utp(path: Path) -> ValidatedUpload:
    return validate_upload(UploadPurpose.UTP, path.name, path.read_bytes())


def _upload_program(path: Path) -> ValidatedUpload:
    return ValidatedUpload(UploadPurpose.PROGRAM, path.name, path.read_bytes(), None)


def test_unknown_program_requires_confirmation_before_plan() -> None:
    program = _program_data(_item("Лепка"))

    assert needs_structure_confirmation(plan=None, program=program) is True
    draft = draft_structure_rows(program=program)
    assert draft[0]["topic"] == ""
    assert "Лепка" not in {row["topic"] for row in draft}
    ledger = draft_source_ledger(program)
    assert ledger[0]["title"] == "Лепка"
    assert ledger[0]["disposition"] == DISPOSITION_UNRESOLVED


def test_unknown_program_becomes_confirmed_study_plan() -> None:
    scope = structure_confirmation_scope(
        program_name="unknown.docx",
        program_digest="abc",
        utp_name=None,
        utp_digest=None,
        study_year=1,
    )
    confirmation = confirm_program_structure(
        rows=(_manual_row(),),
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope=scope,
    )

    assert isinstance(confirmation.plan, ConfirmedStudyPlan)
    assert confirmation.plan.source == "manual"
    assert confirmation.plan.topics[0].title == "Лепка"
    review = confirmation.match_reviews[("1", "Лепка", "Лепка")]
    assert review["decision"] == "USER_CONFIRMED"
    assert confirmation.program_items[0].content == (
        "Знакомство с материалом.\nЛепка из глины."
    )
    original = _program_data(_item("Старое"))
    overlaid = overlay_confirmed_program(original, confirmation.program_items)
    assert confirmation.program_items[0] in overlaid.content_items
    assert any(item.title == "Старое" for item in overlaid.content_items)


def test_empty_fields_are_rejected() -> None:
    scope = "scope"
    with pytest.raises(ProgramStructureConfirmationError, match="тему"):
        confirm_program_structure(
            rows=({**_manual_row(), "topic": ""},),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope=scope,
        )
    with pytest.raises(ProgramStructureConfirmationError, match="часы"):
        confirm_program_structure(
            rows=({**_manual_row(), "theory_hours": ""},),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope=scope,
        )
    empty_source = (
        {
            **_manual_row(),
            "theory_content": "",
            "practice_content": "",
        },
    )
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=empty_source,
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope=scope,
        )


def test_wrong_hours_are_rejected() -> None:
    with pytest.raises(ProgramStructureConfirmationError, match="не совпадают"):
        confirm_program_structure(
            rows=(_manual_row(theory="1", practice="1"),),
            study_year=1,
            study_weeks=1,
            hours_per_week="3",
            scope="scope",
        )
    with pytest.raises(ProgramStructureConfirmationError, match="больше нуля"):
        confirm_program_structure(
            rows=(_manual_row(theory="0", practice="0"),),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="scope",
        )


def test_multiple_embedded_utp_require_confirmation() -> None:
    payload = _program_docx_with_utp_years(1, 2)
    candidates = embedded_utp_candidates(payload, "program.docx")
    program = parse_program(payload, "program.docx", study_year=1)

    assert len(candidates) > 1
    assert needs_structure_confirmation(
        plan=None,
        program=program,
        embedded_utp_count=len(candidates),
        has_external_utp=False,
        study_year=None,
    )
    dummy_plan = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=({"topic": "Тема 1", "total": "2", "theory": "1", "practice": "1"},),
        study_weeks=1,
        hours_per_week="2",
    )
    proven = ContentMatch(
        dummy_plan.topics[0],
        _item("Тема 1"),
        MatchStatus.EXACT,
        1.0,
    )
    assert needs_structure_confirmation(
        plan=dummy_plan,
        program=_program_data(_item("Тема 1")),
        embedded_utp_count=len(candidates),
        has_external_utp=True,
        study_year=1,
        matches=(proven,),
    ) is False


def test_confirmation_resets_when_file_or_year_or_table_changes() -> None:
    rows = (_manual_row(),)
    first = structure_confirmation_scope(
        program_name="a.docx",
        program_digest=file_digest(b"program-a"),
        utp_name=None,
        utp_digest=None,
        study_year=1,
        rows=rows,
    )
    confirmation = confirm_program_structure(
        rows=rows,
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope=first,
    )
    changed_year = structure_confirmation_scope(
        program_name="a.docx",
        program_digest=file_digest(b"program-a"),
        utp_name=None,
        utp_digest=None,
        study_year=2,
        rows=rows,
    )
    changed_file = structure_confirmation_scope(
        program_name="b.docx",
        program_digest=file_digest(b"program-b"),
        utp_name=None,
        utp_digest=None,
        study_year=1,
        rows=rows,
    )
    changed_table = structure_confirmation_scope(
        program_name="a.docx",
        program_digest=file_digest(b"program-a"),
        utp_name=None,
        utp_digest=None,
        study_year=1,
        rows=({**rows[0], "topic": "Другая тема"},),
    )
    assert confirmation_is_current(confirmation, first)
    assert not confirmation_is_current(confirmation, changed_year)
    assert not confirmation_is_current(confirmation, changed_file)
    assert not confirmation_is_current(confirmation, changed_table)


def test_p0_gate_is_not_bypassed_after_structure_confirmation() -> None:
    confirmation = confirm_program_structure(
        rows=(_manual_row(),),
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope="scope",
    )
    source = CalendarContentRow(
        week_number=1,
        date_range="01–07.09",
        month="Сентябрь",
        section="Лепка",
        topic_number="1",
        topic_title="Лепка",
        source_topic_title="Лепка",
        theory_hours=1,
        practice_hours=1,
        total_hours=2,
        match_status=MatchStatus.USER_CONFIRMED,
        program_section="Лепка",
        program_topic="Лепка",
        program_content_full=confirmation.program_items[0].content,
        program_content_preview="Лепка из глины.",
        source_program_name="Неизвестная программа",
        source_utp_name="confirmed",
    )
    row = LessonContentV2Row(
        source=source,
        theory_text="Знакомство с материалом.",
        practice_text="Лепка из глины.",
        lesson_type="практическое занятие",
        planned_result="",
        assessment_method="наблюдение",
        action="лепит",
        object="глина",
        conditions="",
        warnings=(),
        clause_coverage=(),
        clause_roles=(),
    )
    assert EMPTY_RESULT in production_readiness_codes(row)


@pytest.mark.parametrize(
    ("utp_needles", "program_needles", "study_year"),
    (
        (("утп", "ключ", "2"), ("програм", "ключ"), 2),
    ),
)
def test_familiar_kits_stay_on_auto_path(
    utp_needles: tuple[str, ...],
    program_needles: tuple[str, ...],
    study_year: int,
) -> None:
    utp_path = _corpus(*utp_needles)
    program_path = _corpus(*program_needles)
    if utp_path is None or program_path is None:
        pytest.skip(f"Нет комплекта {' '.join(utp_needles)}")
    program = parse_program(
        program_path.read_bytes(),
        program_path.name,
        study_year=study_year,
    )
    plan = resolve_utp(
        _upload_utp(utp_path),
        _upload_program(program_path),
        program_study_year=study_year,
    )
    matches = match_utp_to_program(
        plan.topics,
        program.content_items,
        study_year=study_year,
    )
    assert plan.source == "external_utp"
    assert needs_structure_confirmation(
        plan=plan,
        program=program,
        embedded_utp_count=len(
            embedded_utp_candidates(program_path.read_bytes(), program_path.name)
        ),
        has_external_utp=True,
        study_year=study_year,
        matches=matches,
    ) is False
    assert unmatched_source_topics(plan, matches) == ()


def test_tourists_program_stays_on_auto_path_when_plan_is_confirmed() -> None:
    program_path = _corpus("програм", "турист")
    if program_path is None:
        pytest.skip("Нет программы «Туристы-проводники»")
    program = parse_program(program_path.read_bytes(), program_path.name)
    utp_path = _corpus("утп", "тп", "3")
    plan = None
    if utp_path is not None:
        try:
            plan = resolve_utp(
                _upload_utp(utp_path),
                _upload_program(program_path),
                program_study_year=3,
            )
        except Exception:
            plan = None
    if plan is None:
        assert needs_structure_confirmation(
            plan=None,
            program=program,
            has_external_utp=utp_path is not None,
            study_year=3,
        )
        plan = confirmed_plan_from_manual_rows(
            study_year=1,
            rows=({"topic": program.content_items[0].title, "total": "2", "theory": "1", "practice": "1"},),
            study_weeks=1,
            hours_per_week="2",
        )
        plan = ConfirmedStudyPlan(
            study_year=plan.study_year,
            topics=plan.topics,
            total_hours=plan.total_hours,
            theory_hours=plan.theory_hours,
            practice_hours=plan.practice_hours,
            study_weeks=plan.study_weeks,
            hours_per_week=plan.hours_per_week,
            source="external_utp",
        )
    matches = match_utp_to_program(
        plan.topics,
        program.content_items,
        study_year=plan.study_year,
    )
    assert needs_structure_confirmation(
        plan=plan,
        program=program,
        has_external_utp=True,
        study_year=plan.study_year,
        matches=matches,
    ) is bool(unmatched_source_topics(plan, matches))


def test_source_items_never_disappear() -> None:
    items = (
        _item("Тема A"),
        _item("Тема B"),
        ProgramContentItem(None, "Заголовок", "", "Раздел", 1),
    )
    catalog = catalog_source_items(items)
    partial = (
        _ledger_row(catalog[0], disposition=DISPOSITION_MAPPED, mapped_topic="Тема A"),
    )
    merged = merge_source_ledger(catalog, partial)
    assert [row["item_id"] for row in merged] == [record.item_id for record in catalog]
    counts = source_disposition_counts(merged)
    assert counts["total"] == 3
    assert (
        counts[DISPOSITION_MAPPED]
        + counts[DISPOSITION_EXCLUDED]
        + counts[DISPOSITION_UNRESOLVED]
        == 3
    )
    assert counts[DISPOSITION_UNRESOLVED] == 2
    assert any(row["title"] == "Заголовок" for row in merged)
    confirmation = confirm_program_structure(
        rows=(_manual_row("Тема A"),),
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope="keep-all",
        source_items=items,
        ledger=partial,
    )
    assert [record.item_id for record in confirmation.source_ledger] == [
        record.item_id for record in catalog
    ]


def test_unresolved_source_items_block_confirmation() -> None:
    items = (_item("Лепка"), _item("Роспись"))
    ledger = draft_source_ledger(_program_data(*items))
    assert all(row["disposition"] == DISPOSITION_UNRESOLVED for row in ledger)
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=({**_manual_row(), "theory_content": "", "practice_content": ""},),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="unresolved",
            source_items=items,
            ledger=ledger,
        )


def test_explicit_exclusion_allows_confirmation() -> None:
    items = (
        _item("Лепка"),
        ProgramContentItem(None, "Литература", "Список книг.", "Приложения", 1),
    )
    catalog = catalog_source_items(items)
    ledger = (
        _ledger_row(catalog[0], disposition=DISPOSITION_UNRESOLVED),
        _ledger_row(
            catalog[1],
            disposition=DISPOSITION_EXCLUDED,
            exclusion_reason=EXCLUSION_SERVICE,
        ),
    )
    confirmation = confirm_program_structure(
        rows=(
            {
                **_manual_row(),
                "source_item_id": catalog[0].item_id,
                "excerpt": items[0].content,
                "content_origin": CONTENT_ORIGIN_EXCERPT,
                "theory_content": "",
                "practice_content": "",
            },
        ),
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope="excluded",
        source_items=items,
        ledger=ledger,
    )
    counts = source_disposition_counts(confirmation.source_ledger)
    assert counts[DISPOSITION_MAPPED] == 1
    assert counts[DISPOSITION_EXCLUDED] == 1
    assert counts[DISPOSITION_UNRESOLVED] == 0
    assert confirmation.source_ledger[1].exclusion_reason == EXCLUSION_SERVICE


def test_mapping_or_exclusion_change_resets_confirmation() -> None:
    items = (
        _item("Лепка"),
        ProgramContentItem(None, "Литература", "Список книг.", "Приложения", 1),
    )
    catalog = catalog_source_items(items)
    ledger = (
        _ledger_row(catalog[0], disposition=DISPOSITION_MAPPED, mapped_topic="Лепка"),
        _ledger_row(
            catalog[1],
            disposition=DISPOSITION_EXCLUDED,
            exclusion_reason=EXCLUSION_SERVICE,
        ),
    )
    rows = (_manual_row(),)
    first = structure_confirmation_scope(
        program_name="a.docx",
        program_digest=file_digest(b"program-a"),
        utp_name=None,
        utp_digest=None,
        study_year=1,
        rows=rows,
        ledger=ledger,
    )
    confirmation = confirm_program_structure(
        rows=rows,
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope=first,
        source_items=items,
        ledger=ledger,
    )
    changed_mapping = structure_confirmation_scope(
        program_name="a.docx",
        program_digest=file_digest(b"program-a"),
        utp_name=None,
        utp_digest=None,
        study_year=1,
        rows=rows,
        ledger=(
            _ledger_row(catalog[0], disposition=DISPOSITION_MAPPED, mapped_topic="Другая"),
            ledger[1],
        ),
    )
    changed_exclusion = structure_confirmation_scope(
        program_name="a.docx",
        program_digest=file_digest(b"program-a"),
        utp_name=None,
        utp_digest=None,
        study_year=1,
        rows=rows,
        ledger=(
            ledger[0],
            _ledger_row(
                catalog[1],
                disposition=DISPOSITION_EXCLUDED,
                exclusion_reason=EXCLUSION_HEADING,
            ),
        ),
    )
    assert confirmation_is_current(confirmation, first)
    assert not confirmation_is_current(confirmation, changed_mapping)
    assert not confirmation_is_current(confirmation, changed_exclusion)


def _holdout_year1_schedule():
    holdout = _corpus("природн", "материал")
    if holdout is None:
        pytest.skip("Нет holdout «Работа с природным материалом»")
    program = parse_program(holdout.read_bytes(), holdout.name, study_year=1)
    embedded = embedded_utp_candidates(holdout.read_bytes(), holdout.name)
    selected = select_embedded_utp(embedded, 1)
    assert selected is not None
    topics = schedule_topics_from_candidate(selected)
    assert topics
    return holdout, program, embedded, selected, topics


def test_holdout_uses_explicit_weeks_and_hours_not_topic_count() -> None:
    _holdout, program, embedded, selected, topics = _holdout_year1_schedule()
    draft = draft_structure_rows(
        program=program,
        embedded=embedded,
        study_year=1,
    )
    draft_ledger = draft_source_ledger(
        program,
        study_year=1,
        topics=topics,
    )
    assert needs_structure_confirmation(
        plan=None,
        program=program,
        embedded_utp_count=len(embedded),
        has_external_utp=False,
        study_year=1,
    )
    schedule_titles = tuple(topic.title for topic in topics)
    heading_titles = {"I год обучения", "Особенности работы 1 года обучения"}
    assert [row["topic"] for row in draft] == list(schedule_titles)
    assert set(schedule_titles) != heading_titles
    assert not heading_titles <= {row["topic"] for row in draft}
    assert len(program.content_items) == 45
    assert len(draft_ledger) == 45
    before = source_disposition_counts(draft_ledger)
    assert (
        before[DISPOSITION_MAPPED]
        + before[DISPOSITION_EXCLUDED]
        + before[DISPOSITION_UNRESOLVED]
        == 45
    )
    assert len(draft) == 8
    auto_mapped = [row for row in draft if row["excerpt"]]
    unresolved = unresolved_schedule_rows(draft)
    assert len(auto_mapped) == 6
    assert len(unresolved) == 2
    assert [row["number"] for row in unresolved] == ["7", "8"]
    excerpts = [row["excerpt"] for row in auto_mapped]
    assert len(set(excerpts)) == 6
    source_item = next(
        item
        for item in catalog_source_items(program.content_items)
        if item.item_id == auto_mapped[0]["source_item_id"]
    )
    assert all(excerpt in source_item.content for excerpt in excerpts)
    assert all(excerpt != source_item.content.strip() for excerpt in excerpts)
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=draft,
            study_year=1,
            study_weeks=32,
            hours_per_week="3",
            scope="holdout-unresolved",
            source_items=program.content_items,
            ledger=draft_ledger,
            selected_utp=selected,
            embedded=embedded,
        )
    study_weeks = 32
    hours_per_week = 3
    total_hours = study_weeks * hours_per_week
    plan = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=tuple(
            {
                "topic": row["topic"],
                "total": str(
                    float(row["theory_hours"] or 0) + float(row["practice_hours"] or 0)
                ),
                "theory": row["theory_hours"],
                "practice": row["practice_hours"],
            }
            for row in draft
        ),
        study_weeks=study_weeks,
        hours_per_week=str(hours_per_week),
    )
    schedule = build_schedule(plan)
    assert [topic.title for topic in plan.topics] == list(schedule_titles)
    assert len(plan.topics) == len(topics)
    assert plan.study_weeks == study_weeks
    assert plan.hours_per_week == hours_per_week
    assert plan.total_hours == total_hours
    assert len(plan.topics) != plan.study_weeks
    assert len(schedule.weeks) == study_weeks
    assert sum(element.hours for element in schedule.elements) == total_hours


def test_holdout_mapped_excluded_unresolved_sum_is_45() -> None:
    _holdout, program, _embedded, _selected, topics = _holdout_year1_schedule()
    draft_ledger = draft_source_ledger(program, study_year=1, topics=topics)
    counts = source_disposition_counts(draft_ledger)
    assert len(program.content_items) == 45
    assert counts["total"] == 45
    assert (
        counts[DISPOSITION_MAPPED]
        + counts[DISPOSITION_EXCLUDED]
        + counts[DISPOSITION_UNRESOLVED]
        == 45
    )
    assert counts[DISPOSITION_MAPPED] >= 1
    assert counts[DISPOSITION_EXCLUDED] >= 1
    assert counts[DISPOSITION_UNRESOLVED] >= 1
    years = {row["title"]: row["study_year"] for row in draft_ledger}
    assert years.get("I год обучения") == "1"
    assert years.get("II год обучения") == "2"
    assert years.get("III год обучения") == "3"
    by_title = {row["title"]: row for row in draft_ledger}
    assert by_title["II год обучения"]["disposition"] == DISPOSITION_EXCLUDED
    assert by_title["II год обучения"]["exclusion_reason"] == EXCLUSION_OTHER_YEAR
    assert by_title["III год обучения"]["exclusion_reason"] == EXCLUSION_OTHER_YEAR
    assert by_title["Литература"]["disposition"] == DISPOSITION_EXCLUDED
    assert by_title["Литература"]["exclusion_reason"] == EXCLUSION_SERVICE


def test_headings_do_not_become_schedule_topics() -> None:
    program = _program_data(
        ProgramContentItem(None, "I год обучения", "Текст года.", "I год обучения", 1),
        ProgramContentItem(
            None,
            "Особенности работы 1 года обучения",
            "Особенности.",
            "I год обучения",
            1,
        ),
    )
    draft = draft_structure_rows(program=program)
    ledger = draft_source_ledger(program)
    assert [row["topic"] for row in draft] == [""]
    assert "I год обучения" not in {row["topic"] for row in draft}
    assert {row["title"] for row in ledger} == {
        "I год обучения",
        "Особенности работы 1 года обучения",
    }


def test_selected_utp_rows_are_all_preserved() -> None:
    payload = _program_docx_with_utp_years(1, 2)
    program = parse_program(payload, "program.docx", study_year=1)
    embedded = embedded_utp_candidates(payload, "program.docx")
    selected = select_embedded_utp(embedded, 1)
    expected = schedule_topics_from_candidate(selected)
    draft = tuple(
        {
            **row,
            "practice_content": f"Содержание занятия «{row['topic']}».",
            "theory_content": f"Теория занятия «{row['topic']}».",
            "content_origin": CONTENT_ORIGIN_MANUAL,
        }
        for row in draft_structure_rows(embedded=embedded, study_year=1)
    )
    assert [row["topic"] for row in draft] == [topic.title for topic in expected]
    catalog = catalog_source_items(program.content_items)
    ledger = tuple(
        _ledger_row(record, disposition=DISPOSITION_UNRESOLVED)
        for record in catalog
    )
    weekly = expected[0].hours.total
    weeks = int(sum(topic.hours.total for topic in expected) / weekly)
    confirmation = confirm_program_structure(
        rows=draft,
        study_year=1,
        study_weeks=weeks,
        hours_per_week=str(weekly),
        scope="utp-kept",
        source_items=program.content_items,
        ledger=ledger,
        selected_utp=selected,
        embedded=embedded,
    )
    assert [topic.title for topic in confirmation.plan.topics] == [
        topic.title for topic in expected
    ]
    reduced_weeks = int(sum(topic.hours.total for topic in expected[1:]) / weekly)
    with pytest.raises(ProgramStructureConfirmationError, match="сохранены"):
        confirm_program_structure(
            rows=draft[1:],
            study_year=1,
            study_weeks=reduced_weeks,
            hours_per_week=str(weekly),
            scope="utp-lost",
            source_items=program.content_items,
            ledger=ledger,
            selected_utp=selected,
            embedded=embedded,
        )


def test_content_mapping_is_separate_from_schedule() -> None:
    _holdout, program, embedded, selected, topics = _holdout_year1_schedule()
    draft = draft_structure_rows(program=program, embedded=embedded, study_year=1)
    schedule_titles = {row["topic"] for row in draft}
    content_titles = {item.title for item in program.content_items}
    ledger = draft_source_ledger(program, study_year=1, topics=topics)
    assert all(row["item_id"] for row in ledger)
    assert {row["title"] for row in ledger} == content_titles
    assert len({row["excerpt"] for row in draft if row["excerpt"]}) == 6
    assert {row["topic"] for row in draft} != {row["title"] for row in ledger}
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=draft,
            study_year=1,
            study_weeks=32,
            hours_per_week="3",
            scope="separate",
            source_items=program.content_items,
            ledger=ledger,
            selected_utp=selected,
            embedded=embedded,
        )
    assert schedule_titles == {topic.title for topic in topics}
    assert "I год обучения" not in schedule_titles
    assert "Особенности работы 1 года обучения" not in schedule_titles
    assert "I год обучения" in content_titles
    assert schedule_titles.isdisjoint({"I год обучения", "Особенности работы 1 года обучения"})


def test_schedule_topic_without_source_blocks_confirmation() -> None:
    payload = _program_docx_with_utp_years(1)
    program = parse_program(payload, "program.docx", study_year=1)
    embedded = embedded_utp_candidates(payload, "program.docx")
    selected = select_embedded_utp(embedded, 1)
    draft = draft_structure_rows(embedded=embedded, study_year=1)
    catalog = catalog_source_items(program.content_items)
    ledger = tuple(
        _ledger_row(
            record,
            disposition=DISPOSITION_EXCLUDED,
            exclusion_reason=EXCLUSION_HEADING,
        )
        for record in catalog
    )
    emptied = tuple(
        {**row, "theory_content": "", "practice_content": ""}
        for row in draft
    )
    weekly = schedule_topics_from_candidate(selected)[0].hours.total
    weeks = int(
        sum(topic.hours.total for topic in schedule_topics_from_candidate(selected))
        / weekly
    )
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=emptied,
            study_year=1,
            study_weeks=weeks,
            hours_per_week=str(weekly),
            scope="no-source",
            source_items=program.content_items,
            ledger=ledger,
            selected_utp=selected,
            embedded=embedded,
        )


def test_multi_topic_utp_distributes_across_explicit_weeks() -> None:
    _holdout, program, embedded, selected, topics = _holdout_year1_schedule()
    draft = draft_structure_rows(program=program, embedded=embedded, study_year=1)
    plan = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=tuple(
            {
                "topic": row["topic"],
                "total": str(
                    float(row["theory_hours"] or 0) + float(row["practice_hours"] or 0)
                ),
                "theory": row["theory_hours"],
                "practice": row["practice_hours"],
            }
            for row in draft
        ),
        study_weeks=32,
        hours_per_week="3",
    )
    schedule = build_schedule(plan)
    assert len(plan.topics) == len(topics)
    assert len(plan.topics) > 2
    assert plan.study_weeks == 32
    assert plan.hours_per_week == 3
    assert plan.total_hours == 96
    assert len(schedule.weeks) == 32
    assert sum(element.hours for element in schedule.elements) == 96


def test_utp_topics_and_content_items_never_disappear() -> None:
    _holdout, program, embedded, selected, topics = _holdout_year1_schedule()
    draft = draft_structure_rows(program=program, embedded=embedded, study_year=1)
    catalog = catalog_source_items(program.content_items)
    ledger = draft_source_ledger(program, study_year=1, topics=topics)
    dropped = merge_source_ledger(catalog, ledger[:3])
    assert [row["item_id"] for row in dropped] == [record.item_id for record in catalog]
    assert [row["topic"] for row in draft] == [topic.title for topic in topics]
    assert [row["item_id"] for row in ledger] == [record.item_id for record in catalog]
    assert len(draft) == len(topics)
    assert len(ledger) == 45


def test_auto_plan_with_unmatched_source_opens_confirmation_and_keeps_auto_rows() -> None:
    utp_path = _corpus("утп", "ключ", "1 г")
    program_path = _corpus("програм", "ключ")
    if utp_path is None or program_path is None:
        pytest.skip("Нет комплекта для auto-plan с несопоставленным SOURCE")
    plan = confirmed_plan_from_external_utp(
        parse_utp(utp_path),
        study_year=1,
        source_name=utp_path.name,
    )
    program = parse_program(program_path.read_bytes(), program_path.name, study_year=1)
    matches = match_utp_to_program(plan.topics, program.content_items, study_year=1)
    draft = draft_structure_rows(plan=plan, program=program, matches=matches)
    unmatched = unmatched_source_topics(plan, matches)
    assert unmatched
    assert needs_structure_confirmation(
        plan=plan,
        program=program,
        has_external_utp=True,
        matches=matches,
    )
    assert all(row["match_status"] != MatchStatus.USER_CONFIRMED.value for row in draft)
    before_statuses = {row["topic"]: row["match_status"] for row in draft}
    proven_titles = {
        row["topic"]
        for row in draft
        if row["match_status"] in PROVEN_MATCH_STATUSES
    }
    assert proven_titles
    filled = []
    for row in draft:
        if row["match_status"] in PROVEN_MATCH_STATUSES:
            filled.append(row)
            continue
        filled.append(
            {
                **row,
                "theory_content": row["theory_content"] or "Пользовательское описание теории.",
                "practice_content": row["practice_content"] or "Пользовательское описание практики.",
            }
        )
    confirmation = confirm_program_structure(
        rows=filled,
        study_year=1,
        study_weeks=plan.study_weeks,
        hours_per_week=str(plan.hours_per_week),
        scope="auto-plan",
        existing_plan=plan,
    )
    assert confirmation.plan is plan
    assert confirmation.plan.source == "external_utp"
    unmatched_titles = {topic.title for topic in unmatched}
    assert set(confirmation.match_reviews) == {
        (topic.number, topic.title, topic.parent_section) for topic in unmatched
    }
    assert all(
        review["decision"] == "USER_CONFIRMED"
        for review in confirmation.match_reviews.values()
    )
    overlaid = overlay_confirmed_program(program, confirmation.program_items)
    assert len(overlaid.content_items) >= len(program.content_items)
    after_matches = apply_match_reviews(
        match_utp_to_program(
            plan.topics,
            overlaid.content_items,
            study_year=1,
        ),
        overlaid.content_items,
        confirmation.match_reviews,
    )
    by_title = {match.utp_position.title: match for match in after_matches}
    for title, status in before_statuses.items():
        if title in unmatched_titles:
            assert by_title[title].status is MatchStatus.USER_CONFIRMED
        else:
            assert by_title[title].status.value == status
    schedule = build_schedule(plan)
    before_rows = build_lesson_content_v2(
        build_content_model(schedule, plan, program, utp_path.name)
    )
    after_rows = build_lesson_content_v2(
        build_content_model(
            schedule,
            plan,
            overlaid,
            utp_path.name,
            match_reviews=confirmation.match_reviews,
        )
    )
    assert any(SOURCE_NOT_MATCHED in production_readiness_codes(row) for row in before_rows)
    assert not any(SOURCE_NOT_MATCHED in production_readiness_codes(row) for row in after_rows)


def test_user_confirmed_only_after_explicit_confirm_with_source() -> None:
    proven = {
        **_manual_row("Автотема", "1", "0"),
        "match_status": MatchStatus.EXACT.value,
        "theory_content": "Уже сопоставленный текст.",
        "practice_content": "",
    }
    problem = {
        **_manual_row("Проблемная", "0", "1"),
        "match_status": MatchStatus.NOT_MATCHED.value,
        "theory_content": "",
        "practice_content": "",
    }
    draft = (proven, problem)
    assert all(row["match_status"] != MatchStatus.USER_CONFIRMED.value for row in draft)
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=draft,
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="explicit-empty",
        )
    filled = (
        proven,
        {
            **problem,
            "practice_content": "Наблюдение за привалом.",
            "content_origin": CONTENT_ORIGIN_MANUAL,
        },
    )
    confirmation = confirm_program_structure(
        rows=filled,
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope="explicit",
    )
    assert len(confirmation.match_reviews) == 1
    key = next(iter(confirmation.match_reviews))
    assert key[1] == "Проблемная"
    assert confirmation.match_reviews[key]["decision"] == "USER_CONFIRMED"
    assert confirmation.topic_statuses == (
        MatchStatus.EXACT.value,
        TOPIC_STATUS_USER_CONFIRMED,
    )
    assert all(item.title == "Проблемная" for item in confirmation.program_items)


def test_external_and_manual_sources_remain_available() -> None:
    topic = Topic("1", "Тема", Hours(2, 1, 1), "Тема", True)
    utp = UtpParseResult(
        metadata=UtpMetadata(
            study_year="1 год обучения",
            hours_per_year=2,
            hours_per_week=2,
            study_weeks=1,
        ),
        sections=(Section("1", "Тема", Hours(2, 1, 1), True),),
        topics=(topic,),
        table_totals=Hours(2, 1, 1),
    )
    external = confirmed_plan_from_external_utp(utp, study_year=1)
    manual = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=({"topic": "Тема", "total": "2", "theory": "1", "practice": "1"},),
        study_weeks=1,
        hours_per_week="2",
    )
    assert external.source == "external_utp"
    assert manual.source == "manual"
    proven = ContentMatch(external.topics[0], _item("Тема"), MatchStatus.EXACT, 1.0)
    assert needs_structure_confirmation(
        plan=external,
        program=_program_data(_item("Тема")),
        has_external_utp=True,
        matches=(proven,),
    ) is False


def test_broadcast_mapping_is_forbidden() -> None:
    items = (_item("Общий блок", "Полный текст года обучения."),)
    catalog = catalog_source_items(items)
    ledger = (
        _ledger_row(
            catalog[0],
            disposition=DISPOSITION_MAPPED,
            mapped_topic="Вводное занятие; Аппликация",
        ),
    )
    with pytest.raises(ProgramStructureConfirmationError, match="Broadcast"):
        confirm_program_structure(
            rows=(_manual_row("Вводное занятие"),),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="broadcast-ledger",
            source_items=items,
            ledger=ledger,
        )
    with pytest.raises(ProgramStructureConfirmationError, match="Broadcast"):
        confirm_program_structure(
            rows=(
                {
                    **_manual_row("Вводное занятие"),
                    "mapped_topic": "Вводное занятие | Аппликация",
                    "theory_content": "",
                    "practice_content": "",
                },
            ),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="broadcast-row",
            source_items=items,
        )


def test_each_topic_uses_own_mapping() -> None:
    source = "Часть A про лепку. Часть B про роспись."
    items = (ProgramContentItem("1", "Блок", source, "Блок", 1),)
    catalog = catalog_source_items(items)
    confirmation = confirm_program_structure(
        rows=(
            {
                **_manual_row("Лепка"),
                "source_item_id": catalog[0].item_id,
                "excerpt": "Часть A про лепку.",
                "content_origin": CONTENT_ORIGIN_EXCERPT,
                "theory_content": "",
                "practice_content": "",
            },
            {
                **_manual_row("Роспись"),
                "source_item_id": catalog[0].item_id,
                "excerpt": "Часть B про роспись.",
                "content_origin": CONTENT_ORIGIN_EXCERPT,
                "theory_content": "",
                "practice_content": "",
            },
        ),
        study_year=1,
        study_weeks=2,
        hours_per_week="2",
        scope="own-mapping",
        source_items=items,
    )
    by_title = {item.title: item.content for item in confirmation.program_items}
    assert by_title == {
        "Лепка": "Часть A про лепку.",
        "Роспись": "Часть B про роспись.",
    }
    assert confirmation.topic_statuses == (
        TOPIC_STATUS_USER_CONFIRMED,
        TOPIC_STATUS_USER_CONFIRMED,
    )


def test_one_source_different_excerpts_allowed() -> None:
    source = "Фрагмент теории лепки. Фрагмент практики росписи."
    items = (ProgramContentItem("1", "Блок", source, "Блок", 1),)
    catalog = catalog_source_items(items)
    confirmation = confirm_program_structure(
        rows=(
            {
                **_manual_row("Лепка"),
                "source_item_id": catalog[0].item_id,
                "excerpt": "Фрагмент теории лепки.",
                "content_origin": CONTENT_ORIGIN_EXCERPT,
                "theory_content": "",
                "practice_content": "",
            },
            {
                **_manual_row("Роспись"),
                "source_item_id": catalog[0].item_id,
                "excerpt": "Фрагмент практики росписи.",
                "content_origin": CONTENT_ORIGIN_EXCERPT,
                "theory_content": "",
                "practice_content": "",
            },
        ),
        study_year=1,
        study_weeks=2,
        hours_per_week="2",
        scope="shared-source",
        source_items=items,
    )
    assert [item.content for item in confirmation.program_items] == [
        "Фрагмент теории лепки.",
        "Фрагмент практики росписи.",
    ]
    assert confirmation.source_ledger[0].disposition == DISPOSITION_MAPPED


def test_identical_full_block_for_unrelated_topics_forbidden() -> None:
    source = "Полный блок про лепку и роспись."
    items = (ProgramContentItem("1", "Блок", source, "Блок", 1),)
    catalog = catalog_source_items(items)
    shared = {
        "source_item_id": catalog[0].item_id,
        "excerpt": source,
        "content_origin": CONTENT_ORIGIN_EXCERPT,
        "theory_content": "",
        "practice_content": "",
    }
    with pytest.raises(ProgramStructureConfirmationError, match="full-block"):
        confirm_program_structure(
            rows=(
                {**_manual_row("Лепка"), **shared},
                {**_manual_row("Роспись"), **shared},
            ),
            study_year=1,
            study_weeks=2,
            hours_per_week="2",
            scope="full-block",
            source_items=items,
        )


def test_empty_topic_stays_unresolved() -> None:
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=(
                {
                    **_manual_row("Пустая"),
                    "theory_content": "",
                    "practice_content": "",
                },
            ),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="empty-topic",
        )


def test_user_confirmed_is_per_topic() -> None:
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=(
                {
                    **_manual_row("Подтверждённая"),
                    "content_origin": CONTENT_ORIGIN_MANUAL,
                },
                {
                    **_manual_row("Незаполненная"),
                    "theory_content": "",
                    "practice_content": "",
                },
            ),
            study_year=1,
            study_weeks=2,
            hours_per_week="2",
            scope="per-topic",
        )
    confirmation = confirm_program_structure(
        rows=(
            {
                **_manual_row("Подтверждённая"),
                "content_origin": CONTENT_ORIGIN_MANUAL,
            },
            {
                **_manual_row("Вторая"),
                "content_origin": CONTENT_ORIGIN_MANUAL,
            },
        ),
        study_year=1,
        study_weeks=2,
        hours_per_week="2",
        scope="per-topic-filled",
    )
    assert [key[1] for key in confirmation.match_reviews] == [
        "Подтверждённая",
        "Вторая",
    ]
    assert confirmation.topic_statuses == (
        TOPIC_STATUS_USER_CONFIRMED,
        TOPIC_STATUS_USER_CONFIRMED,
    )


def test_p0_sees_unresolved() -> None:
    confirmation = confirm_program_structure(
        rows=(
            {
                **_manual_row("Подтверждённая"),
                "content_origin": CONTENT_ORIGIN_MANUAL,
            },
        ),
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope="p0-unresolved-confirmed",
    )
    plan = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=(
            {"topic": "Подтверждённая", "total": "2", "theory": "1", "practice": "1"},
            {"topic": "Незаполненная", "total": "2", "theory": "1", "practice": "1"},
        ),
        study_weeks=2,
        hours_per_week="2",
    )
    program = overlay_confirmed_program(None, confirmation.program_items)
    schedule = build_schedule(plan)
    rows = build_lesson_content_v2(
        build_content_model(
            schedule,
            plan,
            program,
            "confirmed",
            match_reviews=confirmation.match_reviews,
        )
    )
    unresolved = [
        row
        for row in rows
        if SOURCE_NOT_MATCHED in production_readiness_codes(row)
    ]
    assert unresolved
    assert any(row.source.topic_title == "Незаполненная" for row in unresolved)


def test_year_scoping_excludes_other_years() -> None:
    program = _program_data(
        ProgramContentItem(
            None,
            "I год обучения",
            "Тема № 1 Лепка из глины руками.",
            "I год обучения",
            1,
        ),
        ProgramContentItem(
            None,
            "Особенности работы 1 года обучения",
            "Репродуктивный метод.",
            "I год обучения",
            1,
        ),
        ProgramContentItem(
            None,
            "II год обучения",
            "Тема № 1 Лепка второго года.",
            "II год обучения",
            1,
        ),
        ProgramContentItem(
            None,
            "III год обучения",
            "Тема № 1 Лепка третьего года.",
            "III год обучения",
            1,
        ),
    )
    topics = (Topic("1", "Лепка", Hours(2, 1, 1), "Лепка", True),)
    ledger = draft_source_ledger(program, study_year=1, topics=topics)
    by_title = {row["title"]: row for row in ledger}
    assert [row["study_year"] for row in ledger] == ["1", "1", "2", "3"]
    assert by_title["I год обучения"]["disposition"] == DISPOSITION_MAPPED
    assert by_title["II год обучения"]["disposition"] == DISPOSITION_EXCLUDED
    assert by_title["II год обучения"]["exclusion_reason"] == EXCLUSION_OTHER_YEAR
    assert by_title["III год обучения"]["exclusion_reason"] == EXCLUSION_OTHER_YEAR
    assert by_title["Особенности работы 1 года обучения"]["disposition"] == (
        DISPOSITION_UNRESOLVED
    )


def test_unique_topic_number_maps_one_to_one_not_broadcast() -> None:
    content = (
        "Тема № 1 Введение в материал и правила.\n"
        "Тема № 2 Лепка фигур из глины."
    )
    program = _program_data(
        ProgramContentItem(None, "I год обучения", content, "I год обучения", 1)
    )
    plan = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=(
            {"topic": "Введение", "total": "1", "theory": "1", "practice": "0"},
            {"topic": "Лепка фигур", "total": "1", "theory": "0", "practice": "1"},
        ),
        study_weeks=1,
        hours_per_week="2",
    )
    draft = draft_structure_rows(plan=plan, program=program, study_year=1)
    excerpts = [row["excerpt"] for row in draft]
    assert excerpts[0].startswith("Тема № 1")
    assert excerpts[1].startswith("Тема № 2")
    assert excerpts[0] != excerpts[1]
    assert excerpts[0] != content.strip()
    assert excerpts[1] != content.strip()
    confirmation = confirm_program_structure(
        rows=draft,
        study_year=1,
        study_weeks=1,
        hours_per_week="2",
        scope="one-to-one",
        source_items=program.content_items,
        ledger=draft_source_ledger(program, study_year=1, topics=plan.topics),
    )
    by_title = {item.title: item.content for item in confirmation.program_items}
    assert by_title["Введение"].startswith("Тема № 1")
    assert by_title["Лепка фигур"].startswith("Тема № 2")


def test_proven_service_and_other_year_are_excluded() -> None:
    program = _program_data(
        ProgramContentItem(
            None,
            "I год обучения",
            "Тема № 1 Введение в материал.",
            "I год обучения",
            1,
        ),
        ProgramContentItem(
            None,
            "Описание условий реализации программы",
            "Кабинет и материалы.",
            "Описание условий реализации программы",
            1,
        ),
        ProgramContentItem(
            None,
            "Методы обучения:",
            "Словесный и наглядный.",
            "Описание условий реализации программы",
            1,
        ),
        ProgramContentItem(
            None,
            "В результате прохождения программы II года обучения учащиеся должны знать:",
            "Правила второго года.",
            "В результате прохождения программы II года обучения учащиеся должны знать:",
            1,
        ),
        ProgramContentItem(None, "Литература", "Список книг.", "Литература", 1),
        ProgramContentItem(
            None,
            "Учебно-тематический план I год обучения",
            "",
            "Учебно-тематический план I год обучения",
            1,
        ),
    )
    topics = (Topic("1", "Введение", Hours(2, 1, 1), "Введение", True),)
    ledger = draft_source_ledger(program, study_year=1, topics=topics)
    by_title = {row["title"]: row for row in ledger}
    assert by_title["Описание условий реализации программы"]["exclusion_reason"] == (
        EXCLUSION_SERVICE
    )
    assert by_title["Методы обучения:"]["exclusion_reason"] == EXCLUSION_SERVICE
    assert by_title["Литература"]["exclusion_reason"] == EXCLUSION_SERVICE
    assert by_title["Учебно-тематический план I год обучения"]["exclusion_reason"] == (
        EXCLUSION_SERVICE
    )
    assert by_title[
        "В результате прохождения программы II года обучения учащиеся должны знать:"
    ]["exclusion_reason"] == EXCLUSION_OTHER_YEAR
    assert source_disposition_counts(ledger)["total"] == 6


def test_duplicate_topic_number_stays_unresolved() -> None:
    content = (
        "Тема № 1 Первый фрагмент введения.\n"
        "Тема № 1 Повтор того же номера."
    )
    program = _program_data(
        ProgramContentItem(None, "I год обучения", content, "I год обучения", 1)
    )
    plan = confirmed_plan_from_manual_rows(
        study_year=1,
        rows=({"topic": "Введение", "total": "2", "theory": "1", "practice": "1"},),
        study_weeks=1,
        hours_per_week="2",
    )
    draft = draft_structure_rows(plan=plan, program=program, study_year=1)
    assert draft[0]["excerpt"] == ""
    assert draft[0]["topic_status"] == TOPIC_STATUS_UNRESOLVED
    ledger = draft_source_ledger(program, study_year=1, topics=plan.topics)
    assert ledger[0]["disposition"] == DISPOSITION_UNRESOLVED
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=draft,
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="ambiguous",
            source_items=program.content_items,
            ledger=ledger,
        )


def test_ui_unknown_program_without_utp_asks_to_confirm_structure() -> None:
    from streamlit.testing.v1 import AppTest

    from test_ui import (
        APP_PATH,
        _check_button,
        _disputed_program_docx,
        _upload_bytes,
    )

    app = AppTest.from_file(str(APP_PATH), default_timeout=30).run()
    _upload_bytes(app, 0, "program-synthetic.docx", _disputed_program_docx())
    app.run()
    year = next(
        item for item in app.number_input if item.label == "Год обучения по программе"
    )
    year.set_value(1).run()
    weeks = next(
        item
        for item in app.number_input
        if item.label == "Количество учебных недель"
    )
    weeks.set_value(1).run()
    weekly = next(
        item
        for item in app.text_input
        if item.label == "Количество часов в неделю"
    )
    weekly.set_value("2").run()
    _check_button(app).click().run()
    assert not app.exception
    assert any(button.label == "Подтвердить структуру" for button in app.button)
    assert app.session_state.get("structure_confirmation_pending") is True
    markdown = " ".join(str(item.value) for item in app.markdown)
    captions = " ".join(str(getattr(item, "value", item)) for item in app.caption)
    visible = f"{markdown} {captions}"
    assert "mapped" in visible
    assert "excluded" in visible
    assert "unresolved" in visible


def test_overlay_keeps_topic_identity_and_hours() -> None:
    draft = (
        {
            "schedule_id": "sch-0001",
            "number": "7",
            "topic": "Итоговое занятие",
            "theory_hours": "2",
            "practice_hours": "0",
            "theory_content": "",
            "practice_content": "",
            "source_item_id": "",
            "excerpt": "",
            "content_origin": "",
            "topic_status": TOPIC_STATUS_UNRESOLVED,
            "source": "program",
            "match_status": MatchStatus.NOT_MATCHED.value,
        },
    )
    edited = (
        {
            **draft[0],
            "topic": "Другое название",
            "number": "99",
            "theory_hours": "8",
            "practice_hours": "4",
            "theory_content": "Подведение итогов года.",
            "content_origin": CONTENT_ORIGIN_MANUAL,
        },
    )
    merged = overlay_unresolved_topic_edits(draft, edited)
    assert merged[0]["topic"] == "Итоговое занятие"
    assert merged[0]["number"] == "7"
    assert merged[0]["theory_hours"] == "2"
    assert merged[0]["practice_hours"] == "0"
    assert merged[0]["theory_content"] == "Подведение итогов года."


def test_empty_required_theory_slot_blocks_confirmation() -> None:
    with pytest.raises(ProgramStructureConfirmationError, match="нерешённые темы"):
        confirm_program_structure(
            rows=(
                {
                    **_manual_row(),
                    "theory_content": "",
                    "practice_content": "Практика заполнена.",
                    "content_origin": CONTENT_ORIGIN_MANUAL,
                },
            ),
            study_year=1,
            study_weeks=1,
            hours_per_week="2",
            scope="empty-theory-slot",
        )


def _open_structure_cards():
    from streamlit.testing.v1 import AppTest

    from test_ui import APP_PATH, _check_button, _upload_bytes

    app = AppTest.from_file(str(APP_PATH), default_timeout=30).run()
    _upload_bytes(app, 0, "cards-program.docx", _card_program_docx())
    app.run()
    year = next(
        item for item in app.number_input if item.label == "Год обучения по программе"
    )
    year.set_value(1).run()
    weeks = next(
        item
        for item in app.number_input
        if item.label == "Количество учебных недель"
    )
    weeks.set_value(1).run()
    weekly = next(
        item
        for item in app.text_input
        if item.label == "Количество часов в неделю"
    )
    weekly.set_value("5").run()
    _check_button(app).click().run()
    assert not app.exception
    assert app.session_state.get("structure_confirmation_pending") is True
    return app


def _visible_confirmation_text(app) -> str:
    markdown = " ".join(str(item.value) for item in app.markdown)
    captions = " ".join(str(getattr(item, "value", item)) for item in app.caption)
    return f"{markdown} {captions}"


def test_ui_unresolved_topics_are_read_only_cards() -> None:
    app = _open_structure_cards()
    visible = _visible_confirmation_text(app)
    assert "Итоговое занятие" in visible
    assert "Конкурсы и выставки" in visible
    assert visible.count("Тема №") >= 2
    assert "Теория: 2 ч · Практика: 0 ч" in visible
    assert "Теория: 0 ч · Практика: 2 ч" in visible
    draft = tuple(app.session_state.get("structure_confirmation_draft") or ())
    unresolved = unresolved_schedule_rows(draft)
    assert len(draft) == 3
    assert len(unresolved) == 2
    assert not any(item.label == "Тема расписания" for item in app.text_input)
    assert STRUCTURE_THEORY_CONTENT_LABEL not in [
        item.label for item in app.text_input
    ]
    assert STRUCTURE_THEORY_CONTENT_LABEL in [item.label for item in app.text_area]
    assert STRUCTURE_PRACTICE_CONTENT_LABEL in [item.label for item in app.text_area]
    assert sum(
        1 for item in app.text_area if item.label == STRUCTURE_THEORY_CONTENT_LABEL
    ) == 1
    assert sum(
        1 for item in app.text_area if item.label == STRUCTURE_PRACTICE_CONTENT_LABEL
    ) == 1
    draft = tuple(app.session_state.get("structure_confirmation_draft") or ())
    assert [row["topic"] for row in draft] == [
        "Введение",
        "Итоговое занятие",
        "Конкурсы и выставки",
    ]


def test_ui_empty_card_fields_block_confirmation() -> None:
    app = _open_structure_cards()
    assert app.session_state.get("structure_confirmation_result") is None
    next(button for button in app.button if button.label == STRUCTURE_CONFIRM_BUTTON).click().run()
    assert app.session_state.get("structure_confirmation_pending") is True
    assert app.session_state.get("structure_confirmation_result") is None
    assert any("нерешённые темы" in item.value for item in app.error)


def test_ui_manual_cards_confirm_without_renaming_topics() -> None:
    app = _open_structure_cards()
    draft = tuple(app.session_state.get("structure_confirmation_draft") or ())
    assert all(
        row.get("match_status") != MatchStatus.USER_CONFIRMED.value for row in draft
    )
    theory = next(
        item for item in app.text_area if item.label == STRUCTURE_THEORY_CONTENT_LABEL
    )
    practice = next(
        item for item in app.text_area if item.label == STRUCTURE_PRACTICE_CONTENT_LABEL
    )
    theory.set_value("Подведение итогов первого года.").run()
    practice.set_value("Подготовка конкурсных работ.").run()
    assert app.session_state.get("structure_confirmation_result") is None
    assert all(
        row.get("match_status") != MatchStatus.USER_CONFIRMED.value
        for row in tuple(app.session_state.get("structure_confirmation_draft") or ())
    )
    next(button for button in app.button if button.label == STRUCTURE_CONFIRM_BUTTON).click().run()
    assert not app.exception
    result = app.session_state.get("structure_confirmation_result")
    assert result is not None
    assert [topic.title for topic in result.plan.topics] == [
        "Введение",
        "Итоговое занятие",
        "Конкурсы и выставки",
    ]
    assert sum(
        status == TOPIC_STATUS_USER_CONFIRMED for status in result.topic_statuses
    ) == 3


def test_ui_source_excerpt_must_be_exact_substring() -> None:
    app = _open_structure_cards()
    radios = [item for item in app.radio if item.label == "Как заполнить содержание"]
    assert len(radios) == 2
    radios[0].set_value(STRUCTURE_ORIGIN_EXCERPT_LABEL).run()
    source = next(
        item for item in app.selectbox if item.label == STRUCTURE_SOURCE_ITEM_LABEL
    )
    ledger = tuple(app.session_state.get("structure_confirmation_draft_ledger") or ())
    mapped = next(row for row in ledger if row["disposition"] == DISPOSITION_MAPPED)
    source.set_value(mapped["item_id"]).run()
    excerpt = next(item for item in app.text_area if item.label == STRUCTURE_EXCERPT_LABEL)
    excerpt.set_value("этого фрагмента нет в SOURCE").run()
    practice = next(
        item for item in app.text_area if item.label == STRUCTURE_PRACTICE_CONTENT_LABEL
    )
    practice.set_value("Подготовка конкурсных работ.").run()
    next(button for button in app.button if button.label == STRUCTURE_CONFIRM_BUTTON).click().run()
    assert app.session_state.get("structure_confirmation_result") is None
    assert any("подстрокой" in item.value for item in app.error)
    excerpt.set_value("Знакомство с инструментами.").run()
    next(button for button in app.button if button.label == STRUCTURE_CONFIRM_BUTTON).click().run()
    result = app.session_state.get("structure_confirmation_result")
    assert result is not None
    assert [topic.title for topic in result.plan.topics][1] == "Итоговое занятие"
