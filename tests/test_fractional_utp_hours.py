from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document
import pytest

from calendar_pedagoga.content_generation import build_content_model
from calendar_pedagoga.parsing import Hours, parse_utp
from calendar_pedagoga.scheduling import build_schedule


REFERENCES = Path(__file__).resolve().parents[1] / "references"


def _fractional_utp_docx(separator: str) -> bytes:
    document = Document()
    document.add_paragraph("Учебно-тематический план")
    table = document.add_table(rows=6, cols=5)
    headers = ("№", "Наименование темы", "Всего", "Теория", "Практика")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    rows = (
        ("1", "Вводное занятие", f"1{separator}5", f"1{separator}5", "0"),
        ("2", "Основной раздел", f"19{separator}5", f"4{separator}5", "15"),
        ("3", "Резерв", "3", "0", "3"),
        ("", "Итого", "24", "6", "18"),
    )
    for index, values in enumerate(rows, start=1):
        for cell, value in zip(table.rows[index].cells, values, strict=True):
            cell.text = value
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("separator", (",", "."))
def test_fractional_utp_hours_are_parsed_without_loss(separator: str) -> None:
    result = parse_utp(_fractional_utp_docx(separator))

    assert result.topics[0].hours == Hours(
        Decimal("1.5"), Decimal("1.5"), 0
    )
    assert result.topics[1].hours == Hours(
        Decimal("19.5"), Decimal("4.5"), 15
    )
    assert result.table_totals == Hours(24, 6, 18)
    assert result.warnings == ()


def test_fractional_utp_hours_survive_schedule_and_content_model() -> None:
    parsed = parse_utp(_fractional_utp_docx(","))
    utp = replace(
        parsed,
        metadata=replace(
            parsed.metadata,
            hours_per_year=24,
            hours_per_week=3,
            study_weeks=8,
        ),
    )

    schedule = build_schedule(utp)
    content = build_content_model(schedule, utp, None, "fractional.docx")

    assert sum(element.hours for element in schedule.elements) == 24
    assert any(element.hours == Decimal("1.5") for element in schedule.elements)
    assert sum(row.total_hours for row in content) == 24
    assert any(
        part.theory_hours == Decimal("1.5")
        for row in content
        for part in row.week_parts
    )
    assert all(row.total_hours == 3 for row in content)


def test_integer_utp_hours_keep_integer_representation() -> None:
    result = parse_utp(REFERENCES / "УТП КЛЮЧ 2 г. 2ч.docx")

    assert result.table_totals == Hours(72, 22, 50)
    assert all(
        isinstance(value, int)
        for topic in result.topics
        for value in (topic.hours.total, topic.hours.theory, topic.hours.practice)
    )
