from io import BytesIO
from types import SimpleNamespace

from docx import Document

from calendar_pedagoga.content_engine_v2 import derive_fields_v2
from calendar_pedagoga.content_generation import (
    _assign_section_blocks,
    _ordered_compact_bindings,
    _section_blocks_from_item,
)
from calendar_pedagoga.matching import ContentMatch, MatchStatus
from calendar_pedagoga.parsing import Hours, Topic
from calendar_pedagoga.program_parsing import ProgramContentItem, infer_study_year_number


def test_section_blocks_follow_source_order_and_keep_core_activity() -> None:
    item = ProgramContentItem(
        "3",
        "Общая физическая подготовка",
        "\n".join(
            (
                "Тема 1. Комплекс упражнений для разминки",
                "Практика",
                "Изучение комплекса упражнений для разминки. Общая разминка.",
                "Тема 2. Упражнения на развитие силовых качеств",
                "Практика",
                "Изучение техники выполнения силовых упражнений. Круговая тренировка.",
                "Тема 3. Упражнения на координацию движений",
                "Практика",
                "Изучение техники выполнения упражнений на координацию. Игра.",
            )
        ),
        "Общая физическая подготовка",
    )

    blocks = _section_blocks_from_item(item)

    assert [block.title for block in blocks] == [
        "Комплекс упражнений для разминки",
        "Упражнения на развитие силовых качеств",
        "Упражнения на координацию движений",
    ]
    assert [block.practice for block in blocks] == [
        "Изучение комплекса упражнений для разминки. Общая разминка.",
        "Изучение техники выполнения силовых упражнений. Круговая тренировка.",
        "Изучение техники выполнения упражнений на координацию. Игра.",
    ]


def test_section_blocks_choose_one_grounded_representative_per_week() -> None:
    items = tuple(
        ProgramContentItem(
            None,
            title,
            f"Практика\n{activity}. Дополнительная общая нагрузка.",
            "Раздел",
        )
        for title, activity in (
            ("Разминка", "Выполнение разминки"),
            ("Сила", "Выполнение силовых упражнений"),
            ("Координация", "Выполнение упражнений на координацию"),
            ("Быстрота", "Выполнение упражнений на быстроту"),
        )
    )
    blocks = tuple(
        block for item in items for block in _section_blocks_from_item(item)
    )

    assigned = _assign_section_blocks(blocks, "practice", appearances=2)

    assert len(assigned) == 2
    text = "\n".join(item.content for item in assigned)
    activities = (
        "Выполнение разминки.",
        "Выполнение силовых упражнений.",
        "Выполнение упражнений на координацию.",
        "Выполнение упражнений на быстроту.",
    )
    for activity in activities:
        assert activity in text
    positions = [text.index(activity) for activity in activities]
    assert positions == sorted(positions)


def test_explicit_activity_form_wins_within_a_dense_week_slot() -> None:
    blocks = tuple(
        _section_blocks_from_item(
            ProgramContentItem(None, title, f"Практика\n{activity}.", "Раздел")
        )[0]
        for title, activity in (
            ("Осанка", "Выполнение упражнений для осанки"),
            ("Корсет", "Выполнение упражнений для мышц спины"),
            ("Внимание", "Игры на развитие внимания"),
        )
    )

    assigned = _assign_section_blocks(blocks, "practice", appearances=1)

    content = assigned[0].content
    assert "Выполнение упражнений для осанки." in content
    assert "Выполнение упражнений для мышц спины." in content
    assert "Игры на развитие внимания." in content


def test_named_diagnostic_prefix_becomes_grounded_week_topic() -> None:
    block = _section_blocks_from_item(
        ProgramContentItem(
            None,
            "Цели программы",
            (
                "Практика\n"
                "Входная диагностика: сдача нормативов по общей физической подготовке."
            ),
            "Вводное занятие",
        )
    )[0]

    assigned = _assign_section_blocks((block,), "practice", appearances=1)

    assert assigned[0].title == "Входная диагностика"


def test_download_copy_suffix_is_not_study_year() -> None:
    assert (
        infer_study_year_number("УТП Скалолазание (2).docx год обучения")
        is None
    )
    assert infer_study_year_number("2 год обучения") == 2
    assert infer_study_year_number("УТП КЛЮЧ 2 г. 2ч.docx") == 2


def test_compact_unmatched_section_uses_only_fully_anchored_source_order() -> None:
    topics = tuple(
        Topic(str(index), title, Hours(2, 0, 2), title, True)
        for index, title in enumerate(("Вводное", "Раздел 2", "Раздел 3"), 1)
    )
    roots = tuple(
        ProgramContentItem(str(index), title, f"Содержание {index}", title)
        for index, title in enumerate(("Цели программы", "Раздел 2", "Раздел 3"), 1)
    )
    matches = {
        (topic.number, topic.title, topic.parent_section): ContentMatch(
            topic,
            roots[index] if index else None,
            MatchStatus.EXACT if index else MatchStatus.NOT_MATCHED,
            1.0 if index else 0.0,
        )
        for index, topic in enumerate(topics)
    }

    resolved = _ordered_compact_bindings(
        SimpleNamespace(
            topics=topics,
            metadata=SimpleNamespace(study_weeks=9, hours_per_week=2),
            table_totals=Hours(6, 0, 6),
        ),
        SimpleNamespace(content_items=roots),
        matches,
        None,
    )

    assert resolved[(topics[0].number, topics[0].title, topics[0].parent_section)] is roots[0]


def test_detailed_utp_is_not_reclassified_as_compact_sections() -> None:
    topics = tuple(
        Topic(str(index), f"Тема {index}", Hours(2, 0, 2), f"Тема {index}", True)
        for index in range(1, 14)
    )
    utp = SimpleNamespace(
        topics=topics,
        metadata=SimpleNamespace(study_weeks=36, hours_per_week=2),
        table_totals=Hours(72, 0, 72),
    )

    assert not _ordered_compact_bindings(
        utp,
        SimpleNamespace(content_items=()),
        {},
        None,
    )


def test_control_inflects_accusative_adjective_noun_phrase() -> None:
    result = derive_fields_v2(
        topic_title="Страховка",
        theory_text="",
        practice_text="Отработка гимнастической страховки напарника",
        program_content="Отработка гимнастической страховки напарника",
        theory_hours=0,
        practice_hours=2,
    )

    assert result.planned_result == "Отрабатывает гимнастическую страховку напарника."
    assert result.assessment_method == (
        "педагогическое наблюдение за отработкой "
        "гимнастической страховки напарника"
    )


def test_broken_genitive_plural_conversion_falls_back_to_week_topic() -> None:
    from calendar_pedagoga.content_engine_v2 import build_lesson_content_v2
    from calendar_pedagoga.content_generation import CalendarContentRow
    from calendar_pedagoga.matching import MatchStatus

    row = CalendarContentRow(
        week_number=1,
        date_range="01.09–07.09",
        month="Сентябрь",
        topic_number="5",
        topic_title="Основы скалолазания",
        source_topic_title="Основы скалолазания",
        section="Основы скалолазания",
        theory_hours=0,
        practice_hours=2,
        total_hours=2,
        match_status=MatchStatus.EXACT,
        program_section="Основы скалолазания",
        program_topic="Действия страховщика",
        program_content_full=(
            "Практика\n"
            "Отработка действий страховщика в зависимости от момента лазанья."
        ),
        program_content_preview="Отработка действий страховщика.",
        source_program_name="program.docx",
        source_utp_name="utp.docx",
    )

    result = build_lesson_content_v2((row,))[0]

    assert result.planned_result == (
        "Выполняет практическое задание по теме „Действия страховщика“."
    )
    assert "действий страховщик" not in result.planned_result.casefold()


def test_different_objects_keep_only_grounded_practice_action() -> None:
    result = derive_fields_v2(
        topic_title="Взаимодействие спортсмена и страховщика",
        theory_text="",
        practice_text=(
            "Изучение голосовых команд и отработка взаимодействия спортсмена "
            "и страховщика."
        ),
        program_content=(
            "Изучение голосовых команд и отработка взаимодействия спортсмена "
            "и страховщика."
        ),
        theory_hours=0,
        practice_hours=2,
    )

    assert result.lesson_type == "практикум по страховке"
    low = result.planned_result.casefold()
    control = result.assessment_method.casefold()
    source = result.practice_text.casefold()
    assert "голосовых команд" in source
    assert "взаимодейств" in source
    assert "взаимодейств" in low
    assert "страховщик" in low
    assert "взаимодейств" in control
    assert "безошибочн" not in low
    assert "свободно применяет" not in low
    if "команд" not in low:
        assert any("NEEDS_REVIEW" in warning for warning in result.warnings)


def test_exercise_control_inflects_deverbal_conjunct() -> None:
    result = derive_fields_v2(
        topic_title="Силовая подготовка",
        theory_text="",
        practice_text="Выполнение упражнений с грузом и удержание отягощений.",
        program_content="Выполнение упражнений с грузом и удержание отягощений.",
        theory_hours=0,
        practice_hours=2,
    )

    assert result.assessment_method == (
        "проверка выполнения упражнений с грузом и удержания отягощений"
    )


def test_theory_adjective_noun_and_paired_action_control_are_grammatical() -> None:
    theory = derive_fields_v2(
        topic_title="Тактика",
        theory_text="Индивидуальная тактика и сильные стороны спортсмена.",
        practice_text="",
        program_content="Индивидуальная тактика и сильные стороны спортсмена.",
        theory_hours=2,
        practice_hours=0,
    )
    practice = derive_fields_v2(
        topic_title="Технические приёмы",
        theory_text="",
        practice_text=(
            "Изучение и отработка основных технических приёмов: "
            "диагональный шаг, накат."
        ),
        program_content=(
            "Изучение и отработка основных технических приёмов: "
            "диагональный шаг, накат."
        ),
        theory_hours=0,
        practice_hours=2,
    )

    assert theory.planned_result == (
        "Характеризует индивидуальную тактику и сильные стороны спортсмена."
    )
    assert "и и" not in practice.assessment_method
    assert "основных технических приёмов" in practice.assessment_method


def test_plural_soft_noun_uses_dative_yam_in_oral_control() -> None:
    result = derive_fields_v2(
        topic_title="Цели программы",
        theory_text="Цели и задачи программы.",
        practice_text="",
        program_content="Цели и задачи программы.",
        theory_hours=2,
        practice_hours=0,
    )

    assert result.assessment_method == "устный опрос по целям и задачам программы"


def test_substantivized_knowledge_head_keeps_grounded_theory_object() -> None:
    result = derive_fields_v2(
        topic_title="Понятие здорового образа жизни",
        theory_text="Составляющие здорового образа жизни.",
        practice_text="",
        program_content="Составляющие здорового образа жизни.",
        theory_hours=2,
        practice_hours=0,
    )

    assert result.planned_result == (
        "Характеризует составляющие здорового образа жизни."
    )
    assert result.assessment_method == (
        "устный опрос по составляющим здорового образа жизни"
    )


def test_named_game_quiz_produces_participation_triad() -> None:
    result = derive_fields_v2(
        topic_title="Спорт и здоровый образ жизни",
        theory_text="",
        practice_text="Игра-викторина «Вода – основа жизни».",
        program_content="Игра-викторина «Вода – основа жизни».",
        theory_hours=0,
        practice_hours=2,
    )

    assert result.lesson_type == "викторина"
    assert result.planned_result == (
        "Участвует в игре-викторине «Вода – основа жизни»."
    )
    assert result.assessment_method == "викторина"


def test_docx_preview_keeps_each_logical_week_together(monkeypatch) -> None:
    from calendar_pedagoga import docx_generation as generation
    from calendar_pedagoga import docx_qa

    def populate(document, *args, **kwargs):
        table = document.add_table(rows=4, cols=8)
        table.rows[2].cells[0].text = "Сентябрь"
        table.rows[2].cells[1].text = "1"
        table.rows[3].cells[0].text = "Сентябрь"
        table.rows[3].cells[1].text = "2"
        return table, generation._columns_for_table(table), ("Сентябрь", "Сентябрь")

    def measure(content, *, total_rows):
        table = Document(BytesIO(content)).tables[0]
        assert total_rows == 2
        assert all(
            row._tr.xpath("./w:trPr/w:cantSplit") for row in table.rows[2:]
        )
        return tuple(
            docx_qa.DataRowPageLayout(
                docx_qa.DataRowPageSpan(1, 1, True),
                (docx_qa.DataRowPageSegment(1, tuple(cell.text for cell in row.cells)),),
            )
            for row in table.rows[2:]
        )

    monkeypatch.setattr(generation, "_load_template", lambda template: Document())
    monkeypatch.setattr(generation, "_populate_calendar_table", populate)
    monkeypatch.setattr(docx_qa, "detect_data_row_page_layout", measure)
    monkeypatch.setattr(generation, "_build_segmented_document", lambda *args, **kwargs: b"docx")

    output = generation.generate_calendar_docx(
        SimpleNamespace(),
        (SimpleNamespace(), SimpleNamespace()),
        SimpleNamespace(uses_organization_template=False),
        "2026–2027",
    )

    assert output == b"docx"


def test_docx_uses_assigned_week_content_even_for_single_occurrence() -> None:
    from calendar_pedagoga import docx_generation as generation
    from calendar_pedagoga.content_generation import WeekTopicPart
    from calendar_pedagoga.matching import MatchStatus

    part = WeekTopicPart(
        topic_number="3",
        topic_title="Общая физическая подготовка",
        section="Общая физическая подготовка",
        theory_hours=0,
        practice_hours=2,
        match_status=MatchStatus.EXACT,
        program_section="Общая физическая подготовка",
        program_topic="Упражнения на координацию",
        program_content_full=(
            "Практика\nВыполнение упражнений на координацию."
        ),
        weekly_content_assigned=True,
    )
    lesson = SimpleNamespace(
        source=SimpleNamespace(source=SimpleNamespace(week_parts=(part,))),
        planned_result="Выполняет упражнения на координацию.",
    )
    key = generation._content_occurrence_key(part)

    _theory, practice = generation._topic_cells_for_lesson(
        lesson,
        {generation._topic_part_key(part): "3"},
        topic_counts={key: 1},
        topic_occurrences={},
    )

    assert "Выполнение упражнений на координацию" in practice
    assert practice != "3. Общая физическая подготовка (2)"


def test_docx_falls_back_to_grounded_week_topic_when_clause_is_not_safe() -> None:
    from calendar_pedagoga import docx_generation as generation
    from calendar_pedagoga.content_generation import WeekTopicPart
    from calendar_pedagoga.matching import MatchStatus

    part = WeekTopicPart(
        topic_number="5",
        topic_title="Основы скалолазания",
        section="Основы скалолазания",
        theory_hours=0,
        practice_hours=2,
        match_status=MatchStatus.EXACT,
        program_section="Основы скалолазания",
        program_topic="Техника лазания по активам",
        program_content_full="Практика\nЛазание учебных трасс.",
        weekly_content_assigned=True,
    )
    lesson = SimpleNamespace(
        source=SimpleNamespace(source=SimpleNamespace(week_parts=(part,))),
        planned_result=(
            "Выполняет практическое задание по теме „Техника лазания по активам“."
        ),
    )
    key = generation._content_occurrence_key(part)

    _theory, practice = generation._topic_cells_for_lesson(
        lesson,
        {generation._topic_part_key(part): "5"},
        topic_counts={key: 1},
        topic_occurrences={},
    )

    assert "Лазание учебных трасс" in practice
    assert practice != "Техника лазания по активам (2)"


def test_docx_uses_source_clause_that_matches_the_week_result() -> None:
    from calendar_pedagoga import docx_generation as generation
    from calendar_pedagoga.content_generation import WeekTopicPart
    from calendar_pedagoga.matching import MatchStatus

    part = WeekTopicPart(
        topic_number="3",
        topic_title="Общая физическая подготовка",
        section="Общая физическая подготовка",
        theory_hours=0,
        practice_hours=2,
        match_status=MatchStatus.EXACT,
        program_section="Общая физическая подготовка",
        program_topic="Упражнения для улучшения осанки",
        program_content_full=(
            "Практика\n"
            "Изучение упражнений для улучшения осанки.\n"
            "Игры на развитие внимания."
        ),
        weekly_content_assigned=True,
    )
    lesson = SimpleNamespace(
        source=SimpleNamespace(source=SimpleNamespace(week_parts=(part,))),
        planned_result="Участвует в играх на развитие внимания.",
    )
    key = generation._content_occurrence_key(part)

    _theory, practice = generation._topic_cells_for_lesson(
        lesson,
        {generation._topic_part_key(part): "3"},
        topic_counts={key: 1},
        topic_occurrences={},
    )

    assert "Изучение упражнений для улучшения осанки" in practice
    assert "Игры на развитие внимания" in practice
def test_training_nominal_produces_observable_grounded_triad() -> None:
    result = derive_fields_v2(
        topic_title="Постановка ног",
        theory_text="",
        practice_text="Тренировка постановки ног при помощи игр на равновесие.",
        program_content="Тренировка постановки ног при помощи игр на равновесие.",
        theory_hours=0,
        practice_hours=2,
    )

    assert result.lesson_type == "учебно-тренировочное занятие"
    assert result.planned_result == (
        "Отрабатывает постановку ног при помощи игр на равновесие."
    )
    assert result.assessment_method == (
        "педагогическое наблюдение за отработкой постановки ног "
        "при помощи игр на равновесие"
    )


def test_practical_study_of_technique_becomes_observable_rehearsal() -> None:
    result = derive_fields_v2(
        topic_title="Техника лазания",
        theory_text="",
        practice_text="Изучение техники лазания в упор и распор.",
        program_content="Изучение техники лазания в упор и распор.",
        theory_hours=0,
        practice_hours=2,
    )

    assert result.lesson_type == "учебно-тренировочное занятие"
    assert result.planned_result == "Отрабатывает технику лазания в упор и распор."
    assert result.assessment_method == "педагогическое наблюдение за техникой лазания"


def test_nominal_climbing_actions_keep_concrete_source_activity() -> None:
    climbing = derive_fields_v2(
        topic_title="Лазание по активам",
        theory_text="",
        practice_text="Лазание по трассам с активными зацепами.",
        program_content="Лазание по трассам с активными зацепами.",
        theory_hours=0,
        practice_hours=2,
    )
    hangs = derive_fields_v2(
        topic_title="Висы",
        theory_text="",
        practice_text="Висы на зацепках, планках, турнике.",
        program_content="Висы на зацепках, планках, турнике.",
        theory_hours=0,
        practice_hours=2,
    )

    assert climbing.lesson_type == "учебно-тренировочное занятие"
    assert climbing.planned_result == "Выполняет лазание по трассам с активными зацепами."
    assert climbing.assessment_method == (
        "педагогическое наблюдение за выполнением лазания по трассам с активными зацепами"
    )
    assert hangs.lesson_type == "учебно-тренировочное занятие"
    assert hangs.planned_result == "Выполняет висы на зацепках, планках, турнике."
    assert hangs.assessment_method == (
        "педагогическое наблюдение за выполнением висов на зацепках, планках, турнике"
    )

def test_compound_diary_product_gets_grounded_product_type() -> None:
    result = derive_fields_v2(
        topic_title="Дневник тренировок",
        theory_text="",
        practice_text="Составление и ведение дневника тренировок.",
        program_content="Составление и ведение дневника тренировок.",
        theory_hours=0,
        practice_hours=2,
    )

    assert result.lesson_type == "проектно-практическое занятие"
    low = result.planned_result.casefold()
    control = result.assessment_method.casefold()
    assert "составляет" in low
    assert "ведёт" in low or "ведет" in low
    assert "дневник" in low
    assert "дневник" in control
