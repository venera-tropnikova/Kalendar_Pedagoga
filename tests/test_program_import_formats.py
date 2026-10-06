"""Format-neutral fixtures: no corpus filenames or subject-specific runtime rules."""
from io import BytesIO
from pathlib import Path
import subprocess

from docx import Document
import pymupdf
import pytest

from calendar_pedagoga.program_extraction import read_docx, ProgramReadError
from calendar_pedagoga.program_parsing import parse_program, parse_program_docx, convert_legacy_doc, find_libreoffice, LegacyDocUnsupportedError
from calendar_pedagoga.upload_validation import validate_upload, UploadPurpose, UploadValidationError


def word(lines):
    doc = Document()
    for text in lines:
        doc.add_paragraph(text)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pdf(pages):
    doc = pymupdf.open()
    font = next((p for p in [Path("C:/Windows/Fonts/arial.ttf"),
                            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf")] if p.is_file()), None)
    assert font, "A Cyrillic font is required for the test fixture"
    for lines in pages:
        page = doc.new_page()
        if lines is None:
            image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 30, 30), False)
            image.clear_with(255)
            page.insert_image(page.rect, stream=image.tobytes("png"))
        else:
            page.insert_font(fontname="fixture", fontfile=str(font))
            for index, text in enumerate(lines):
                page.insert_text((30, 40 + index * 22), text, fontname="fixture", fontsize=11)
    return doc.tobytes()


@pytest.mark.parametrize("heading", [
    "Содержание программы", "Содержание курса", "Содержание учебного материала",
    "Содержательная часть", "Содержание программы «Произвольное название»",
])
def test_unbolded_headings_and_original_order(heading):
    result = parse_program_docx(word([heading, "1. Раздел А", "Первое действие.", "2. Раздел Б", "Второе действие."]))
    assert [x.title for x in result.content_items] == ["Раздел А", "Раздел Б"]
    assert [x.content for x in result.content_items] == ["Первое действие.", "Второе действие."]
    assert all(x.source_locations for x in result.content_items)


def test_content_in_table_and_nested_table_no_duplicate_cells():
    doc = Document()
    doc.add_paragraph("Содержание курса")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Тема"
    table.cell(0, 1).text = "Содержание"
    table.cell(1, 0).text = "Альфа"
    table.cell(1, 1).text = "Первый фрагмент."
    table.cell(2, 0).text = "Бета"
    table.cell(2, 1).text = "Повтор."
    nested = table.cell(2, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).merge(nested.cell(0, 1)).text = "Вложенный фрагмент."
    doc.add_paragraph("После таблицы.")
    stream = BytesIO(); doc.save(stream)
    extracted = read_docx(stream.getvalue())
    texts = [b.text for b in extracted.paragraphs]
    assert texts.count("Вложенный фрагмент.") == 1
    assert texts.index("Первый фрагмент.") < texts.index("Бета") < texts.index("После таблицы.")
    result = parse_program_docx(stream.getvalue())
    assert [x.title for x in result.content_items] == ["Альфа", "Бета"]
    assert result.content_items[1].content == "Повтор.\nВложенный фрагмент.\nПосле таблицы."


def test_same_text_distinct_locations():
    result = read_docx(word(["Повтор.", "Повтор."]))
    assert result.paragraphs[0].text == result.paragraphs[1].text
    assert result.paragraphs[0].location != result.paragraphs[1].location


def test_multiple_years_require_choice_and_do_not_mix():
    data = word(["Содержание курса", "Первый год обучения", "1. Одна тема", "Источник А.",
                 "Второй год обучения", "1. Одна тема", "Источник Б."])
    result = parse_program_docx(data)
    assert result.import_needs_confirmation and not result.content_items
    assert {r.year for r in result.content_candidates} == {1, 2}
    second = parse_program_docx(data, study_year=2)
    assert [x.content for x in second.content_items] == ["Источник Б."]
    assert second.content_items[0].study_year == 2
    with pytest.raises(ProgramReadError):
        parse_program_docx(data, study_year=1, content_selection=second.selected_content)


def test_result_mention_does_not_switch_year():
    data = word(["Содержание программы 2-го года обучения", "1. Раздел",
                 "Результаты первого года обсуждаются на занятии.", "Продолжение второго года."])
    result = parse_program_docx(data, study_year=2)
    assert result.content_items[0].study_year == 2
    assert "Продолжение" in result.content_items[0].content


def test_inferred_after_plan_requires_confirmation_and_keeps_original_text():
    doc = Document()
    doc.add_paragraph("Цель программы: обучение.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Название темы"
    table.cell(0, 1).text = "Часы"
    table.cell(1, 0).text = "Произвольный блок"
    table.cell(1, 1).text = "4"
    doc.add_paragraph("Произвольный блок(4)")
    doc.add_paragraph("Описание занятия.")
    doc.add_paragraph("Практическая работа.")
    doc.add_paragraph("Другое действие.")
    doc.add_paragraph("Должны знать:")
    doc.add_paragraph("Ожидаемый результат.")
    stream = BytesIO(); doc.save(stream)
    data = stream.getvalue()
    result = validate_upload(UploadPurpose.PROGRAM, "renamed.docx", data).parsed
    assert result.import_needs_confirmation and not result.content_items
    confirmed = parse_program_docx(data, content_selection=result.content_candidates[0].id)
    assert len(confirmed.content_items) == 1
    assert confirmed.content_items[0].content == "Описание занятия.\nПрактическая работа.\nДругое действие."
    assert parse_program(data, "different.docx") == parse_program(data, "renamed.docx")


def test_text_pdf_keeps_page_provenance():
    data = pdf([["Программа «Пример»", "Содержание курса", "1. Первый раздел", "Первый текст."],
                ["2. Второй раздел", "Второй текст."]])
    result = validate_upload(UploadPurpose.PROGRAM, "anonymous.pdf", data).parsed
    assert [x.content for x in result.content_items] == ["Первый текст.", "Второй текст."]
    assert {b.page for b in result.source_blocks} == {1, 2}
    assert "страница 2" in result.content_items[1].source_locations[0]


@pytest.mark.parametrize("pages,missing", [
    ([None], "1"),
    ([["Содержание курса", "1. Раздел", "Текст."], None], "2"),
    ([None, ["Текстовая страница."], None], "1, 3"),
])
def test_scanned_pdf_is_not_silently_partially_accepted(pages, missing):
    with pytest.raises(UploadValidationError, match="страницы " + missing) as exc:
        validate_upload(UploadPurpose.PROGRAM, "scan.pdf", pdf(pages))
    assert "OCR на русском" in str(exc.value)


def test_no_content_has_specific_location():
    with pytest.raises(UploadValidationError, match="Последнее прочитанное место: document/body"):
        validate_upload(UploadPurpose.PROGRAM, "empty-structure.docx", word(["Отдельная заметка."]))


def test_doc_conversion_timeout_is_explained(monkeypatch):
    monkeypatch.setattr("calendar_pedagoga.program_parsing.find_libreoffice", lambda: Path("soffice"))
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("soffice", 120)
    monkeypatch.setattr("calendar_pedagoga.program_parsing.subprocess.run", timeout)
    with pytest.raises(LegacyDocUnsupportedError, match="120 с"):
        convert_legacy_doc(b"legacy")


@pytest.mark.skipif(find_libreoffice() is None, reason="LibreOffice required for real DOC conversion")
def test_real_doc_conversion_roundtrip(tmp_path):
    original = word(["Содержание курса", "1. Раздел", "Исходное действие."])
    (tmp_path / "source.docx").write_bytes(original)
    converter = find_libreoffice()
    output = subprocess.run([str(converter), "--headless", "-env:UserInstallation=" + (tmp_path/"profile").as_uri(),
                             "--convert-to", "doc:MS Word 97", "--outdir", str(tmp_path),
                             str(tmp_path/"source.docx")], capture_output=True, timeout=120)
    assert output.returncode == 0 and (tmp_path/"source.doc").is_file()
    result = validate_upload(UploadPurpose.PROGRAM, "renamed.doc", (tmp_path/"source.doc").read_bytes()).parsed
    assert result.content_items[0].content == "Исходное действие."


def test_plain_short_headings_are_proposals_not_guesses():
    data = word(["Содержание курса", "Обычный заголовок", "Описание первого занятия.",
                 "Другой заголовок", "Описание второго занятия."])
    result = parse_program_docx(data)
    assert result.import_needs_confirmation and not result.content_items
    region = result.content_candidates[0]
    assert region.heading_indices
    confirmed = parse_program_docx(data, content_selection=region.id)
    assert [x.title for x in confirmed.content_items] == ["Обычный заголовок", "Другой заголовок"]


def test_duplicate_boundaries_are_not_resolved_automatically():
    result = parse_program_docx(word(["Содержание курса", "1. Раздел", "Текст А.",
                                     "Содержание курса", "1. Раздел", "Текст Б."]))
    assert len(result.content_candidates) == 2
    assert not result.content_items and result.import_needs_confirmation


def test_missing_headings_report_exact_source_location():
    with pytest.raises(ProgramReadError, match="document/body"):
        parse_program_docx(word(["Содержание курса", "Текст без структуры."]))


def test_pdf_without_plan_requests_separate_plan_not_docx_error():
    from calendar_pedagoga.resolve_utp import resolve_utp, UtpResolutionError
    uploaded = validate_upload(UploadPurpose.PROGRAM, "text.pdf",
                               pdf([["Содержание курса", "1. Раздел", "Описание занятия."]]))
    with pytest.raises(UtpResolutionError, match="Загрузите УТП отдельным файлом"):
        resolve_utp(None, uploaded)


def test_results_between_years_do_not_hide_later_content():
    data = word(["Содержание курса", "Первый год обучения", "1. Раздел", "Действие А.",
                 "Ожидаемые результаты", "Знания первого года.",
                 "Второй год обучения", "1. Раздел", "Действие Б.",
                 "Ожидаемые результаты", "Знания второго года."])
    result = parse_program_docx(data)
    assert {r.year for r in result.content_candidates} == {1, 2}
    assert parse_program_docx(data, study_year=1).content_items[0].content == "Действие А."
    assert parse_program_docx(data, study_year=2).content_items[0].content == "Действие Б."


def test_repeated_year_after_results_does_not_reopen_content():
    data = word(["\u0421\u043e\u0434\u0435\u0440\u0436\u0430\u043d\u0438\u0435 \u043a\u0443\u0440\u0441\u0430", "1 \u0433\u043e\u0434 \u043e\u0431\u0443\u0447\u0435\u043d\u0438\u044f", "1. Section", "Teaching text.",
                 "\u041e\u0436\u0438\u0434\u0430\u0435\u043c\u044b\u0435 \u0440\u0435\u0437\u0443\u043b\u044c\u0442\u0430\u0442\u044b", "Outcomes.",
                 "1 \u0433\u043e\u0434 \u043e\u0431\u0443\u0447\u0435\u043d\u0438\u044f", "1. Assessment", "Question."])
    result = parse_program_docx(data)
    assert not result.import_needs_confirmation
    assert len(result.content_candidates) == 1
    assert result.content_items[0].content == "Teaching text."


def russian_scan():
    native = pdf([["Содержание курса", "1. Раздел", "Изучение природы родного края."]])
    with pymupdf.open(stream=native, filetype="pdf") as source:
        raster = source[0].get_pixmap(dpi=300)
        with pymupdf.open() as scanned:
            page = scanned.new_page(width=source[0].rect.width, height=source[0].rect.height)
            page.insert_image(page.rect, stream=raster.tobytes("png"))
            return scanned.tobytes(deflate=True)


def test_real_russian_ocr_scan():
    from calendar_pedagoga.program_extraction import read_pdf
    data = russian_scan()
    with pymupdf.open(stream=data, filetype="pdf") as source:
        assert not source[0].get_text().strip()
    result = read_pdf(data)
    assert "Изучение природы родного края." in " ".join(b.text for b in result.paragraphs)
    assert all(b.page == 1 and "OCR" in b.location for b in result.paragraphs)
    program = validate_upload(UploadPurpose.PROGRAM, "anonymous.pdf", data).parsed
    assert program.content_items[0].content == "Изучение природы родного края."


def test_ocr_unavailable_reports_page_and_reason(monkeypatch):
    import calendar_pedagoga.program_extraction as extraction
    def unavailable():
        raise ProgramReadError("OCR недоступен: русская модель отсутствует")
    monkeypatch.setattr(extraction, "ocr_tessdata", unavailable)
    with pytest.raises(UploadValidationError, match="страницы 1.*русская модель"):
        validate_upload(UploadPurpose.PROGRAM, "anonymous.pdf", russian_scan())


def test_text_pdf_does_not_call_ocr(monkeypatch):
    import calendar_pedagoga.program_extraction as extraction
    def unexpected(page):
        pytest.fail("Textual PDF must not use OCR")
    monkeypatch.setattr(extraction, "_ocr_page", unexpected)
    result = extraction.read_pdf(pdf([["Содержание курса", "1. Раздел", "Исходный текст."]]))
    assert all("OCR" not in b.location for b in result.paragraphs)
