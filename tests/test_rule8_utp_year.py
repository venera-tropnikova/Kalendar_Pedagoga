from io import BytesIO
from pathlib import Path

from docx import Document
import pytest

from calendar_pedagoga.content_generation import study_year_for_matching
from calendar_pedagoga.matching import match_utp_to_program
from calendar_pedagoga.parsing import Hours, UtpYearSelectionError, parse_utp
from calendar_pedagoga.program_parsing import infer_study_year_number, parse_program
from calendar_pedagoga.resolve_utp import (
    RECONCILE_LEAD,
    RECONCILE_NOTICE,
    RECONCILE_PASS,
    UtpResolutionError,
    compare_embedded_to_separate,
    resolve_utp,
)
from calendar_pedagoga.upload_validation import UploadPurpose, ValidatedUpload, validate_upload


REFERENCES = Path(__file__).resolve().parents[1] / "references"


def _add_year_table(
    document,
    year: int | None,
    totals: Hours,
    *,
    topic_title: str = "Введение в тему",
    extra_paragraph: str | None = None,
) -> None:
    if year is None:
        document.add_paragraph("Учебно-тематический план")
    else:
        document.add_paragraph(
            f"Учебно-тематический план {year}-го года обучения"
        )
    if extra_paragraph:
        document.add_paragraph(extra_paragraph)
    table = document.add_table(rows=6, cols=5)
    headers = ("№", "Тема", "всего", "теория", "практика")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    third = totals.total - 4
    theory_third = max(totals.theory - 2, 0)
    practice_third = third - theory_third
    rows = (
        ("1", "Раздел первый", str(totals.total), str(totals.theory), str(totals.practice)),
        ("1.1", topic_title, "2", "1", "1"),
        ("1.2", "Практика раздела", "2", "1", "1"),
        ("1.3", "Закрепление материала", str(third), str(theory_third), str(practice_third)),
        ("Итого", "", str(totals.total), str(totals.theory), str(totals.practice)),
    )
    for index, values in enumerate(rows, start=1):
        for cell, value in zip(table.rows[index].cells, values, strict=True):
            cell.text = value


def _multi_year_program(*tables: tuple[int | None, Hours, str]) -> bytes:
    document = Document()
    document.add_paragraph("Дополнительная общеобразовательная программа «ТЕСТ»")
    for year, totals, topic_title in tables:
        _add_year_table(document, year, totals, topic_title=topic_title)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _separate_utp(
    year: int | None,
    totals: Hours,
    *,
    topic_title: str = "Введение в тему",
    filename: str | None = None,
) -> ValidatedUpload:
    document = Document()
    if year is not None:
        document.add_paragraph(f"Год обучения: {year}")
    document.add_paragraph("Количество часов в неделю: 2")
    document.add_paragraph(f"Общее количество часов в год: {totals.total}")
    document.add_paragraph("36 учебных недель")
    _add_year_table(document, year, totals, topic_title=topic_title)
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()
    name = filename or (f"УТП {year} г.docx" if year else "УТП.docx")
    return ValidatedUpload(UploadPurpose.UTP, name, data, parse_utp(data))


def _program_upload(data: bytes, name: str = "program.docx") -> ValidatedUpload:
    return ValidatedUpload(UploadPurpose.PROGRAM, name, data, None)


def test_parse_utp_selects_requested_second_year() -> None:
    data = _multi_year_program(
        (1, Hours(36, 12, 24), "Тема первого года"),
        (2, Hours(72, 24, 48), "Тема второго года"),
    )
    result = parse_utp(data, study_year=2)
    assert result.table_totals == Hours(72, 24, 48)
    assert any(topic.title == "Тема второго года" for topic in result.topics)
    assert infer_study_year_number(result.metadata.study_year) == 2


def test_parse_utp_blocks_when_year_unknown_and_several_tables() -> None:
    data = _multi_year_program(
        (1, Hours(36, 12, 24), "Тема первого года"),
        (2, Hours(72, 24, 48), "Тема второго года"),
    )
    with pytest.raises(UtpYearSelectionError, match="год не выбран"):
        parse_utp(data)


def test_parse_utp_blocks_when_requested_year_missing() -> None:
    data = _multi_year_program(
        (1, Hours(36, 12, 24), "Тема первого года"),
        (2, Hours(72, 24, 48), "Тема второго года"),
    )
    with pytest.raises(UtpYearSelectionError, match="3-го года"):
        parse_utp(data, study_year=3)


def test_resolve_same_year_identical_structure_is_pass() -> None:
    totals = Hours(72, 24, 48)
    program = _program_upload(
        _multi_year_program((2, totals, "Введение в тему"))
    )
    separate = _separate_utp(2, totals)
    result = resolve_utp(separate, program)
    assert result.table_totals == totals
    assert result.topics == separate.parsed.topics
    status, diffs = compare_embedded_to_separate(
        parse_utp(program.content, study_year=2),
        separate.parsed,
    )
    assert status == RECONCILE_PASS
    assert diffs == ()
    assert RECONCILE_LEAD not in result.warnings


def test_resolve_same_year_different_hours_is_notice() -> None:
    program = _program_upload(
        _multi_year_program((2, Hours(144, 48, 96), "Введение в тему"))
    )
    separate = _separate_utp(2, Hours(72, 24, 48))
    result = resolve_utp(separate, program)
    assert result.table_totals == Hours(72, 24, 48)
    assert result.topics == separate.parsed.topics
    assert RECONCILE_LEAD in result.warnings
    assert any("144/48/96" in warning and "72/24/48" in warning for warning in result.warnings)


def test_resolve_same_year_different_topics_is_notice_with_diff() -> None:
    program = _program_upload(
        _multi_year_program((2, Hours(72, 24, 48), "Тема программы"))
    )
    separate = _separate_utp(2, Hours(72, 24, 48), topic_title="Тема отдельного плана")
    result = resolve_utp(separate, program)
    assert result.table_totals == Hours(72, 24, 48)
    assert any(topic.title == "Тема отдельного плана" for topic in result.topics)
    assert RECONCILE_LEAD in result.warnings
    assert any("Тема программы" in warning for warning in result.warnings)
    assert any("Тема отдельного плана" in warning for warning in result.warnings)


def test_key_program_only_blocks_without_requested_year() -> None:
    program_path = REFERENCES / "Программа КЛЮЧ.DOC"
    program = validate_upload(
        UploadPurpose.PROGRAM, program_path.name, program_path.read_bytes()
    )
    with pytest.raises(UtpResolutionError, match="год не выбран"):
        resolve_utp(None, program)


def test_key_embedded_year_two_is_not_first_year_hours() -> None:
    program_path = REFERENCES / "Программа КЛЮЧ.DOC"
    result = parse_utp(program_path.read_bytes(), study_year=2)
    assert result.table_totals != Hours(113, 35, 78)
    assert result.table_totals == Hours(144, 34, 110)
    assert infer_study_year_number(result.metadata.study_year) == 2


def test_key_separate_year_two_stays_source_with_adaptation_notice() -> None:
    utp_path = REFERENCES / "УТП КЛЮЧ 2 г. 2ч.docx"
    program_path = REFERENCES / "Программа КЛЮЧ.DOC"
    validated_utp = validate_upload(UploadPurpose.UTP, utp_path.name, utp_path.read_bytes())
    validated_program = validate_upload(
        UploadPurpose.PROGRAM, program_path.name, program_path.read_bytes()
    )
    result = resolve_utp(validated_utp, validated_program)
    assert result.table_totals == Hours(72, 22, 50)
    assert len(result.topics) == 13
    assert infer_study_year_number(result.metadata.study_year) == 2
    assert RECONCILE_LEAD in result.warnings
    assert any("144" in warning and "72" in warning for warning in result.warnings)
    status, diffs = compare_embedded_to_separate(
        parse_utp(program_path.read_bytes(), study_year=2),
        validated_utp.parsed,
    )
    assert status == RECONCILE_NOTICE
    assert diffs


def test_separate_year_two_does_not_fall_back_to_similar_year_one() -> None:
    similar = Hours(72, 24, 48)
    program = _program_upload(
        _multi_year_program((1, similar, "Введение в тему"))
    )
    separate = _separate_utp(2, similar)
    with pytest.raises(UtpResolutionError, match="противоречат"):
        resolve_utp(separate, program)


def test_tour_guides_explicit_year_three_versus_program_year_one_blocks() -> None:
    utp_path = REFERENCES / "УТП ТП 3г. 2ч.docx"
    program_path = REFERENCES / "Программа ТУРИСТЫ-ПРОВОДНИКИ 1 г.docx"
    validated_utp = validate_upload(UploadPurpose.UTP, utp_path.name, utp_path.read_bytes())
    validated_program = validate_upload(
        UploadPurpose.PROGRAM, program_path.name, program_path.read_bytes()
    )
    with pytest.raises(UtpResolutionError, match="противоречат"):
        resolve_utp(validated_utp, validated_program)


def test_key_separate_and_tour_guides_program_year_conflict_blocks() -> None:
    key_utp = REFERENCES / "УТП КЛЮЧ 2 г. 2ч.docx"
    tp_program = REFERENCES / "Программа ТУРИСТЫ-ПРОВОДНИКИ 1 г.docx"
    validated_utp = validate_upload(UploadPurpose.UTP, key_utp.name, key_utp.read_bytes())
    validated_program = validate_upload(
        UploadPurpose.PROGRAM, tp_program.name, tp_program.read_bytes()
    )
    with pytest.raises(UtpResolutionError, match="противоречат"):
        resolve_utp(validated_utp, validated_program)


def test_parse_program_and_matching_receive_selected_year() -> None:
    program_path = REFERENCES / "Программа КЛЮЧ.DOC"
    utp = parse_utp(program_path.read_bytes(), study_year=2)
    selected = study_year_for_matching(utp)
    assert selected == 2
    year2 = parse_program(program_path.read_bytes(), program_path.name, study_year=2)
    year1 = parse_program(program_path.read_bytes(), program_path.name, study_year=1)
    filtered = parse_program(
        program_path.read_bytes(),
        program_path.name,
        study_year=selected,
    )
    assert [item.title for item in filtered.content_items] == [
        item.title for item in year2.content_items
    ]
    assert [item.title for item in filtered.content_items] != [
        item.title for item in year1.content_items
    ]
    assert all(
        item.study_year in {None, 2} for item in filtered.content_items
    )
    matches = match_utp_to_program(
        utp.topics,
        filtered.content_items,
        study_year=selected,
    )
    assert matches
    assert all(
        match.program_item is None or match.program_item.study_year in {None, 2}
        for match in matches
    )
